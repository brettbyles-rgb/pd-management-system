from __future__ import annotations

import argparse
import csv
import hashlib
import heapq
import json
import math
import re
import shutil
import sqlite3
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook


BUILD_VERSION = "1.1.0"
ROLE_NAMESPACE = uuid.UUID("2a16f2d0-74eb-4d86-a306-6626f518b98a")
CAPABILITY_NAMESPACE = uuid.UUID("ac7f5054-1612-4529-8099-57d0b85e28c5")


@dataclass(frozen=True)
class Paths:
    root: Path
    database_override: Path | None = None

    @property
    def database(self) -> Path:
        return self.database_override or (
            self.root / "output" / "pd-management-bulk-fixed.sqlite3"
        )

    @property
    def workbook(self) -> Path:
        return (
            self.root
            / "data"
            / "source"
            / "original"
            / "activity-assigned-clustered-role-links-revised-high-all.xlsx"
        )

    @property
    def legacy_dir(self) -> Path:
        return self.root / "data" / "source" / "onedrive_exports"

    @property
    def canonical(self) -> Path:
        return self.root / "data" / "source" / "canonical"

    @property
    def reference(self) -> Path:
        return self.root / "data" / "reference"

    @property
    def generated(self) -> Path:
        return self.root / "data" / "generated"

    @property
    def qa(self) -> Path:
        return self.root / "data" / "qa"

    @property
    def docs(self) -> Path:
        return self.root / "docs"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _operational_database_hash(path: Path) -> str:
    """Hash only maintained inputs used by the pathway import.

    The unified file also contains generated pathway tables. Hashing the entire
    SQLite file would make the next release depend on its own previous output.
    """
    queries = {
        "position_descriptions": "SELECT * FROM position_descriptions ORDER BY id",
        "classification_references": "SELECT * FROM classification_references ORDER BY id",
        "active_job_family_entries": "SELECT * FROM active_job_family_entries ORDER BY id",
        "active_pd_job_family_mappings": "SELECT * FROM active_pd_job_family_mappings ORDER BY id",
        "capability_frameworks": "SELECT * FROM capability_frameworks ORDER BY id",
        "capability_definitions": "SELECT * FROM capability_definitions ORDER BY id",
        "pd_capabilities": "SELECT * FROM pd_capabilities ORDER BY id",
        "essential_requirements": (
            "SELECT * FROM pd_list_items WHERE section_name = 'essential_requirements' "
            "ORDER BY id"
        ),
    }
    digest = hashlib.sha256()
    connection = _connect_database(path)
    try:
        for name, query in queries.items():
            digest.update(name.encode("utf-8"))
            cursor = connection.execute(query)
            digest.update("|".join(column[0] for column in cursor.description).encode("utf-8"))
            for row in cursor:
                digest.update(
                    json.dumps(list(row), ensure_ascii=False, separators=(",", ":"), default=str).encode(
                        "utf-8"
                    )
                )
    finally:
        connection.close()
    return digest.hexdigest()


def _write_csv(path: Path, fields: Sequence[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _csv_value(row.get(field)) for field in fields})


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def _stable_role_id(source_db_id: int) -> str:
    return str(uuid.uuid5(ROLE_NAMESPACE, f"tafe-nsw-position-description:{source_db_id}"))


def _stable_capability_id(framework: str, name: str, code: str | None) -> str:
    return str(
        uuid.uuid5(
            CAPABILITY_NAMESPACE,
            f"{framework.strip()}|{(code or '').strip()}|{name.strip()}",
        )
    )


def _normalised_level(raw_level: str) -> float | None:
    text = raw_level.strip()
    public_sector = {
        "Foundational": 1,
        "Intermediate": 2,
        "Adept": 3,
        "Advanced": 4,
        "Highly Advanced": 5,
    }
    if text in public_sector:
        return float(public_sector[text])
    match = re.search(r"\bLevel\s+([1-9])\b", text, re.IGNORECASE)
    if match:
        return float(match.group(1))
    return None


def _match_text(value: str) -> str:
    """Normalise known legacy CSV encoding damage for matching only."""
    text = value.strip()
    if "â" in text or "Ã" in text:
        try:
            text = text.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return re.sub(r"[‐‑‒–—−]", "-", text).casefold()


def _connect_database(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def _existing_role_ids(path: Path) -> dict[int, str]:
    if not path.exists():
        return {}
    return {
        int(row["source_db_id"]): row["role_id"]
        for row in _read_csv(path)
        if row.get("source_db_id") and row.get("role_id")
    }


def import_sources(paths: Paths) -> dict[str, Any]:
    for directory in (paths.canonical, paths.reference, paths.qa):
        directory.mkdir(parents=True, exist_ok=True)
    if not paths.database.exists():
        raise FileNotFoundError(paths.database)
    if not paths.workbook.exists():
        raise FileNotFoundError(paths.workbook)

    connection = _connect_database(paths.database)
    role_rows = [
        dict(row)
        for row in connection.execute(
            """SELECT id, source_filename, role_title, position_description_no,
                      extraction_status, department_agency, division_branch_unit,
                      classification_grade_band, anzsco_code, osca_code, pcat_code,
                      date_of_approval, senior_executive_work_level_standards
               FROM position_descriptions ORDER BY id"""
        )
    ]
    prior_ids = _existing_role_ids(paths.canonical / "roles.csv")
    role_id_by_db_id: dict[int, str] = {}
    canonical_roles: list[dict[str, Any]] = []
    for row in role_rows:
        source_db_id = int(row["id"])
        role_id = prior_ids.get(source_db_id) or _stable_role_id(source_db_id)
        role_id_by_db_id[source_db_id] = role_id
        canonical_roles.append(
            {
                "role_id": role_id,
                "source_db_id": source_db_id,
                "source_filename": row["source_filename"],
                "pd_number": row["position_description_no"],
                "title": row["role_title"],
                "classification_raw": row["classification_grade_band"],
                "extraction_status": row["extraction_status"],
                "department_agency": row["department_agency"],
                "division_branch_unit": row["division_branch_unit"],
                "anzsco_code": row["anzsco_code"],
                "osca_code": row["osca_code"],
                "pcat_code": row["pcat_code"],
                "date_of_approval": row["date_of_approval"],
                "senior_executive_work_level_standards": row[
                    "senior_executive_work_level_standards"
                ],
                "record_status": "active",
            }
        )
    _write_csv(
        paths.canonical / "roles.csv",
        [
            "role_id",
            "source_db_id",
            "source_filename",
            "pd_number",
            "title",
            "classification_raw",
            "extraction_status",
            "department_agency",
            "division_branch_unit",
            "anzsco_code",
            "osca_code",
            "pcat_code",
            "date_of_approval",
            "senior_executive_work_level_standards",
            "record_status",
        ],
        canonical_roles,
    )

    classification_rows = [
        dict(row)
        for row in connection.execute(
            """SELECT id AS classification_id, raw_label, display_label,
                      abbreviation, cohort, seniority_order, enterprise_agreement,
                      active, notes, updated_at
               FROM classification_references ORDER BY id"""
        )
    ]
    _write_csv(
        paths.reference / "classification_aliases.csv",
        [
            "classification_id",
            "raw_label",
            "display_label",
            "abbreviation",
            "cohort",
            "seniority_order",
            "enterprise_agreement",
            "active",
            "notes",
            "updated_at",
        ],
        classification_rows,
    )

    job_family_nodes = [
        dict(row)
        for row in connection.execute(
            """SELECT code, name, level, parent_code, main_definition,
                      supp_definition_1, supp_definition_2, exclusions, sequence
               FROM active_job_family_entries ORDER BY sequence, code"""
        )
    ]
    _write_csv(
        paths.reference / "job_family_nodes.csv",
        [
            "code",
            "name",
            "level",
            "parent_code",
            "main_definition",
            "supp_definition_1",
            "supp_definition_2",
            "exclusions",
            "sequence",
        ],
        job_family_nodes,
    )

    mapping_rows = []
    for row in connection.execute(
        """SELECT position_description_id, pd_id AS legacy_mapping_pd_id,
                  mapping_rank, mapping_code, mapping_name, mapping_level,
                  mapping_validated, validation_notes, framework_code_valid
           FROM active_pd_job_family_mappings
           ORDER BY position_description_id, mapping_rank, id"""
    ):
        db_id = row["position_description_id"]
        mapping_rows.append(
            {
                "role_id": role_id_by_db_id.get(int(db_id)) if db_id else None,
                **dict(row),
            }
        )
    _write_csv(
        paths.canonical / "role_job_family_mappings.csv",
        [
            "role_id",
            "position_description_id",
            "legacy_mapping_pd_id",
            "mapping_rank",
            "mapping_code",
            "mapping_name",
            "mapping_level",
            "mapping_validated",
            "validation_notes",
            "framework_code_valid",
        ],
        mapping_rows,
    )

    framework_rows = [
        dict(row)
        for row in connection.execute(
            "SELECT id AS source_framework_id, framework_name FROM capability_frameworks "
            "ORDER BY id"
        )
    ]
    for row in framework_rows:
        row["framework_id"] = f"framework:{row['source_framework_id']}"
    _write_csv(
        paths.reference / "capability_frameworks.csv",
        ["framework_id", "source_framework_id", "framework_name"],
        framework_rows,
    )

    capability_rows: list[dict[str, Any]] = []
    capability_id_by_source: dict[int, str] = {}
    for row in connection.execute(
        """SELECT cd.id AS source_capability_id, cf.id AS source_framework_id,
                  cf.framework_name,
                  cd.capability_name, cd.capability_code
           FROM capability_definitions cd
           JOIN capability_frameworks cf ON cf.id = cd.framework_id
           ORDER BY cd.id"""
    ):
        capability_id = _stable_capability_id(
            row["framework_name"], row["capability_name"], row["capability_code"]
        )
        capability_id_by_source[int(row["source_capability_id"])] = capability_id
        capability_rows.append(
            {
                "capability_id": capability_id,
                "framework_id": f"framework:{row['source_framework_id']}",
                **dict(row),
            }
        )
    _write_csv(
        paths.reference / "capabilities.csv",
        [
            "capability_id",
            "framework_id",
            "source_capability_id",
            "source_framework_id",
            "framework_name",
            "capability_code",
            "capability_name",
        ],
        capability_rows,
    )

    role_capability_rows = []
    observed_levels: Counter[tuple[str, str]] = Counter()
    capability_roles: set[int] = set()
    focus_roles: set[int] = set()
    for row in connection.execute(
        """SELECT pc.id AS source_assignment_id, pc.position_description_id,
                  pc.capability_definition_id, pc.sequence, pc.capability_type,
                  pc.required_level, pc.capability_group, pc.description,
                  pc.source_text, cf.framework_name, cd.capability_name,
                  cd.capability_code
           FROM pd_capabilities pc
           JOIN capability_definitions cd ON cd.id = pc.capability_definition_id
           JOIN capability_frameworks cf ON cf.id = cd.framework_id
           ORDER BY pc.position_description_id, pc.sequence, pc.id"""
    ):
        db_id = int(row["position_description_id"])
        capability_roles.add(db_id)
        if str(row["capability_type"]).lower() == "focus":
            focus_roles.add(db_id)
        raw_level = str(row["required_level"] or "")
        observed_levels[(row["framework_name"], raw_level)] += 1
        role_capability_rows.append(
            {
                "role_id": role_id_by_db_id[db_id],
                "capability_id": capability_id_by_source[
                    int(row["capability_definition_id"])
                ],
                **dict(row),
                "normalized_level": _normalised_level(raw_level),
            }
        )
    _write_csv(
        paths.canonical / "role_capabilities.csv",
        [
            "role_id",
            "capability_id",
            "source_assignment_id",
            "position_description_id",
            "sequence",
            "capability_type",
            "framework_name",
            "capability_code",
            "capability_name",
            "required_level",
            "normalized_level",
            "capability_group",
            "description",
            "source_text",
        ],
        role_capability_rows,
    )
    level_rules = [
        {
            "framework_name": framework,
            "raw_level": raw_level,
            "normalized_level": _normalised_level(raw_level),
            "assignment_count": count,
            "interpretation_status": (
                "mapped" if _normalised_level(raw_level) is not None else "review_required"
            ),
        }
        for (framework, raw_level), count in sorted(observed_levels.items())
    ]
    _write_csv(
        paths.reference / "capability_level_rules.csv",
        [
            "framework_name",
            "raw_level",
            "normalized_level",
            "assignment_count",
            "interpretation_status",
        ],
        level_rules,
    )

    classification_by_raw = {
        row["raw_label"]: row for row in classification_rows if row["raw_label"]
    }
    applicability_rows = []
    for role in canonical_roles:
        db_id = int(role["source_db_id"])
        classification = classification_by_raw.get(role["classification_raw"])
        if db_id in capability_roles:
            status, reason = "applicable", "capability records present"
        else:
            status, reason = "review_required", "no capability records; confirm missing or not applicable"
        applicability_rows.append(
            {
                "role_id": role["role_id"],
                "applicability_status": status,
                "reason": reason,
                "classification_cohort": classification.get("cohort") if classification else None,
                "has_focus_capabilities": db_id in focus_roles,
            }
        )
    _write_csv(
        paths.canonical / "capability_applicability.csv",
        [
            "role_id",
            "applicability_status",
            "reason",
            "classification_cohort",
            "has_focus_capabilities",
        ],
        applicability_rows,
    )

    requirement_rows = []
    for row in connection.execute(
        """SELECT position_description_id, sequence, item_text
           FROM pd_list_items WHERE section_name = 'essential_requirements'
           ORDER BY position_description_id, sequence, id"""
    ):
        db_id = int(row["position_description_id"])
        requirement_rows.append(
            {
                "role_id": role_id_by_db_id[db_id],
                "requirement_type": "essential_requirement",
                **dict(row),
            }
        )
    _write_csv(
        paths.canonical / "essential_requirements.csv",
        [
            "role_id",
            "position_description_id",
            "sequence",
            "requirement_type",
            "item_text",
        ],
        requirement_rows,
    )

    workbook = load_workbook(paths.workbook, read_only=True, data_only=True)
    link_sheet = workbook["Role Activity Links"]
    link_rows = link_sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(link_rows)]
    activities: dict[str, dict[str, Any]] = {}
    role_activity_rows = []
    unresolved_activity_ids: list[Any] = []
    for values in link_rows:
        row = dict(zip(headers, values))
        if not row.get("Activity ID"):
            continue
        activity_id = str(row["Activity ID"])
        activities.setdefault(
            activity_id,
            {
                "activity_id": activity_id,
                "canonical_match_label": row.get("Canonical Match Label"),
                "canonical_plain_label": row.get("Canonical Plain Label"),
                "discriminating": row.get("Discriminating"),
                "cluster_id": row.get("Cluster ID"),
                "raw_count": row.get("Raw Count"),
                "cluster_member_phrases": row.get("Cluster Member Phrases"),
            },
        )
        try:
            db_id = int(row["PD ID"])
        except (TypeError, ValueError):
            unresolved_activity_ids.append(row.get("PD ID"))
            continue
        role_id = role_id_by_db_id.get(db_id)
        if role_id is None:
            unresolved_activity_ids.append(db_id)
            continue
        role_activity_rows.append(
            {
                "role_id": role_id,
                "position_description_id": db_id,
                "pd_number": row.get("PD Number"),
                "title": row.get("Title"),
                "activity_rank": row.get("Activity Rank"),
                "activity_id": activity_id,
                "assignment_gaps": row.get("Assignment Gaps"),
                "retry_count": row.get("Retry Count"),
                "fallback_used": row.get("Fallback Used"),
            }
        )
    workbook.close()
    _write_csv(
        paths.reference / "activities.csv",
        [
            "activity_id",
            "canonical_match_label",
            "canonical_plain_label",
            "discriminating",
            "cluster_id",
            "raw_count",
            "cluster_member_phrases",
        ],
        sorted(activities.values(), key=lambda row: row["activity_id"]),
    )
    _write_csv(
        paths.canonical / "role_activities.csv",
        [
            "role_id",
            "position_description_id",
            "pd_number",
            "title",
            "activity_rank",
            "activity_id",
            "assignment_gaps",
            "retry_count",
            "fallback_used",
        ],
        role_activity_rows,
    )
    _write_csv(
        paths.reference / "adjacency_rules.csv",
        [
            "rule_id",
            "source_family_code",
            "target_family_code",
            "direction",
            "weight",
            "rationale",
            "status",
        ],
        [],
    )

    operational_database_hash = _operational_database_hash(paths.database)
    source_files = [paths.workbook]
    source_release_hash = hashlib.sha256(
        (
            f"operational-database:{operational_database_hash}|"
            + "|".join(f"{path.name}:{_sha256(path)}" for path in source_files)
        ).encode()
    ).hexdigest()
    release = {
        "source_release": source_release_hash[:16],
        "build_version": BUILD_VERSION,
        "source_files": [
            {
                "path": str(paths.database.relative_to(paths.root)),
                "sha256": operational_database_hash,
                "hash_scope": "maintained pathway input rows",
            },
            *[
            {"path": str(path.relative_to(paths.root)), "sha256": _sha256(path)}
            for path in source_files
            ],
        ],
        "record_counts": {
            "roles": len(canonical_roles),
            "role_capabilities": len(role_capability_rows),
            "role_activities": len(role_activity_rows),
            "activities": len(activities),
            "role_job_family_mappings": len(mapping_rows),
            "essential_requirements": len(requirement_rows),
        },
    }
    _write_json(paths.canonical / "source_release.json", release)
    _write_json(
        paths.canonical / "build_config.json",
        {
            "build_version": BUILD_VERSION,
            "algorithm_version": "provisional-1.0.0",
            "neighbours_per_role": 25,
            "weights": {
                "target_focus_capability_readiness": 0.55,
                "rarity_weighted_activity_similarity": 0.30,
                "same_subfamily": 0.10,
                "classification_proximity": 0.05,
            },
            "activity_rarity_formula": "1 - roles_with_activity / total_active_roles",
            "capability_normalisation": "source raw value retained; interpretation stored separately",
            "capability_not_applicable_policy": "requires explicit per-role status; no cohort is inferred",
            "complete_rebuild": True,
        },
    )
    connection.close()
    return {
        **release,
        "unresolved_activity_workbook_ids": sorted(
            {str(value) for value in unresolved_activity_ids}
        ),
    }


def _role_data(paths: Paths) -> tuple[list[dict[str, str]], dict[str, dict[str, str]]]:
    roles = _read_csv(paths.canonical / "roles.csv")
    return roles, {row["role_id"]: row for row in roles}


def _classification_data(paths: Paths) -> tuple[dict[str, dict[str, str]], dict[str, int]]:
    rows = _read_csv(paths.reference / "classification_aliases.csv")
    by_raw = {row["raw_label"]: row for row in rows}
    ordered = sorted(
        {
            float(row["seniority_order"])
            for row in rows
            if row.get("seniority_order") not in (None, "")
        }
    )
    ranks = {str(value): index for index, value in enumerate(ordered)}
    return by_raw, ranks


def _primary_family_maps(paths: Paths) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]]]:
    nodes = {row["code"]: row for row in _read_csv(paths.reference / "job_family_nodes.csv")}
    mappings = _read_csv(paths.canonical / "role_job_family_mappings.csv")
    primary: dict[str, dict[str, str]] = {}
    for mapping in mappings:
        if mapping["role_id"] and mapping["mapping_rank"] == "1":
            primary[mapping["role_id"]] = mapping
    hierarchy: dict[str, dict[str, str]] = {}
    for role_id, mapping in primary.items():
        chain: list[dict[str, str]] = []
        node = nodes.get(mapping["mapping_code"])
        seen: set[str] = set()
        while node and node["code"] not in seen:
            seen.add(node["code"])
            chain.append(node)
            node = nodes.get(node["parent_code"])
        values = {"job_family": "", "sub_family": "", "specialisation": ""}
        for item in reversed(chain):
            level = item["level"].lower()
            if "job family" in level:
                values["job_family"] = item["name"]
            elif "sub" in level:
                values["sub_family"] = item["name"]
            elif "special" in level:
                values["specialisation"] = item["name"]
        hierarchy[role_id] = values
    return primary, hierarchy


def _legacy_crosswalk(
    paths: Paths,
    roles: list[dict[str, str]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    by_db_id = {int(row["source_db_id"]): row for row in roles}
    by_title_grade: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for role in roles:
        by_title_grade[
            (_match_text(role["title"]), _match_text(role["classification_raw"]))
        ].append(role)
    crosswalk: list[dict[str, Any]] = []
    exceptions: list[dict[str, Any]] = []
    for role in roles:
        crosswalk.append(
            {
                "role_id": role["role_id"],
                "source_system": "pd_management_database",
                "legacy_identifier_type": "position_description_id",
                "legacy_identifier": role["source_db_id"],
                "match_status": "exact",
                "match_basis": "database primary key",
                "pd_number": role["pd_number"],
                "title": role["title"],
                "classification": role["classification_raw"],
            }
        )
    profile_path = paths.legacy_dir / "pd_capability_profile_ac.csv"
    if profile_path.exists():
        legacy_roles: dict[str, dict[str, str]] = {}
        for row in _read_csv(profile_path):
            legacy_roles.setdefault(row["pd_id"], row)
        for legacy_id, legacy in sorted(legacy_roles.items(), key=lambda item: int(item[0])):
            candidates = by_title_grade.get(
                (_match_text(legacy["pd_title"]), _match_text(legacy["pd_grade"])), []
            )
            status = "exact" if len(candidates) == 1 else "unresolved" if not candidates else "ambiguous"
            role_id = candidates[0]["role_id"] if len(candidates) == 1 else ""
            entry = {
                "role_id": role_id,
                "source_system": "legacy_capability_profile_and_neighbour_graph",
                "legacy_identifier_type": "sequential_array_position",
                "legacy_identifier": legacy_id,
                "match_status": status,
                "match_basis": "unique exact title and classification" if role_id else "requires review",
                "pd_number": candidates[0]["pd_number"] if role_id else "",
                "title": legacy["pd_title"],
                "classification": legacy["pd_grade"],
            }
            crosswalk.append(entry)
            if not role_id:
                exceptions.append(
                    {
                        "issue_type": "legacy_role_id_crosswalk",
                        "severity": "error",
                        "role_id": "",
                        "record_key": legacy_id,
                        "description": f"{status} legacy title/classification match: {legacy['pd_title']} | {legacy['pd_grade']}",
                        "required_action": "business owner selects the current role or confirms retirement",
                    }
                )
    rarity_path = paths.legacy_dir / "pd_activity_rarity.csv"
    if rarity_path.exists():
        for legacy_id, legacy in {
            row["pd_id"]: row for row in _read_csv(rarity_path)
        }.items():
            candidate = by_db_id.get(int(legacy_id)) if legacy_id.isdigit() else None
            exact = candidate is not None and candidate["title"] == legacy["pd_title"]
            crosswalk.append(
                {
                    "role_id": candidate["role_id"] if exact else "",
                    "source_system": "legacy_activity_rarity",
                    "legacy_identifier_type": "position_description_id",
                    "legacy_identifier": legacy_id,
                    "match_status": "exact" if exact else "unresolved",
                    "match_basis": "database id and title" if exact else "requires review",
                    "pd_number": candidate["pd_number"] if exact else "",
                    "title": legacy["pd_title"],
                    "classification": candidate["classification_raw"] if exact else "",
                }
            )
    return crosswalk, exceptions


def _generate_activity_rarity(
    paths: Paths,
    roles: list[dict[str, str]],
    role_by_id: dict[str, dict[str, str]],
) -> tuple[list[dict[str, Any]], dict[str, float]]:
    links = _read_csv(paths.canonical / "role_activities.csv")
    activities = {row["activity_id"]: row for row in _read_csv(paths.reference / "activities.csv")}
    role_count = len([row for row in roles if row["record_status"] == "active"])
    role_sets: dict[str, set[str]] = defaultdict(set)
    for link in links:
        role_sets[link["activity_id"]].add(link["role_id"])
    activity_counts = {
        activity_id: len(role_ids) for activity_id, role_ids in role_sets.items()
    }
    shares = {
        activity_id: round(count / role_count, 12)
        for activity_id, count in activity_counts.items()
    }
    rows = []
    for link in links:
        role = role_by_id[link["role_id"]]
        activity = activities[link["activity_id"]]
        share = shares[link["activity_id"]]
        rows.append(
            {
                "role_id": link["role_id"],
                "pd_id": role["source_db_id"],
                "pd_title": role["title"],
                "activity_id": link["activity_id"],
                "activity_label": activity["canonical_plain_label"],
                "roles_with_activity": activity_counts[link["activity_id"]],
                "total_active_roles": role_count,
                "share": share,
                "rarity_weight": round(1.0 - share, 12),
            }
        )
    return rows, shares


def _generate_capability_profile(
    paths: Paths,
    role_by_id: dict[str, dict[str, str]],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, float]]]:
    assignments = _read_csv(paths.canonical / "role_capabilities.csv")
    rows = []
    focus: dict[str, dict[str, float]] = defaultdict(dict)
    for assignment in assignments:
        normalized = (
            float(assignment["normalized_level"])
            if assignment["normalized_level"]
            else None
        )
        role = role_by_id[assignment["role_id"]]
        is_focus = assignment["capability_type"].lower() == "focus"
        rows.append(
            {
                "role_id": assignment["role_id"],
                "pd_id": role["source_db_id"],
                "pd_title": role["title"],
                "pd_grade": role["classification_raw"],
                "framework": assignment["framework_name"],
                "capability_code": assignment["capability_code"],
                "capability": assignment["capability_name"],
                "raw_level": assignment["required_level"],
                "level": normalized,
                "focus": is_focus,
            }
        )
        if is_focus and normalized is not None:
            focus[assignment["role_id"]][assignment["capability_id"]] = normalized
    return rows, focus


def _weighted_jaccard(left: set[str], right: set[str], weights: dict[str, float]) -> float:
    union = left | right
    if not union:
        return 0.0
    intersection = left & right
    denominator = sum(weights.get(item, 1.0) for item in union)
    return sum(weights.get(item, 1.0) for item in intersection) / denominator if denominator else 0.0


def _readiness(source: dict[str, float], target: dict[str, float]) -> float:
    if not target:
        return 0.0
    return sum(min(source.get(key, 0.0) / level, 1.0) for key, level in target.items() if level) / len(target)


def _generate_neighbours(
    paths: Paths,
    roles: list[dict[str, str]],
    role_by_id: dict[str, dict[str, str]],
    focus: dict[str, dict[str, float]],
    activity_shares: dict[str, float],
    hierarchy: dict[str, dict[str, str]],
    classification_by_raw: dict[str, dict[str, str]],
    classification_ranks: dict[str, int],
    neighbours_per_role: int,
    weights: dict[str, float],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    activity_sets: dict[str, set[str]] = defaultdict(set)
    for link in _read_csv(paths.canonical / "role_activities.csv"):
        activity_sets[link["role_id"]].add(link["activity_id"])
    rarity = {key: 1.0 - share for key, share in activity_shares.items()}
    heaps: dict[str, list[tuple[float, str, dict[str, Any]]]] = defaultdict(list)

    def grade_rank(role: dict[str, str]) -> int | None:
        reference = classification_by_raw.get(role["classification_raw"])
        if not reference or not reference.get("seniority_order"):
            return None
        return classification_ranks.get(str(float(reference["seniority_order"])))

    for index, source_role in enumerate(roles):
        source_id = source_role["role_id"]
        source_rank = grade_rank(source_role)
        for target_role in roles[index + 1 :]:
            target_id = target_role["role_id"]
            target_rank = grade_rank(target_role)
            work_similarity = _weighted_jaccard(
                activity_sets[source_id], activity_sets[target_id], rarity
            )
            readiness_forward = _readiness(focus[source_id], focus[target_id])
            readiness_reverse = _readiness(focus[target_id], focus[source_id])
            same_subfamily = bool(
                hierarchy.get(source_id, {}).get("sub_family")
                and hierarchy.get(source_id, {}).get("sub_family")
                == hierarchy.get(target_id, {}).get("sub_family")
            )
            grade_delta = (
                target_rank - source_rank
                if source_rank is not None and target_rank is not None
                else None
            )
            grade_proximity = 0.0 if grade_delta is None else max(0.0, 1.0 - abs(grade_delta) / 4.0)
            for left, right, ready, delta in (
                (source_role, target_role, readiness_forward, grade_delta),
                (
                    target_role,
                    source_role,
                    readiness_reverse,
                    -grade_delta if grade_delta is not None else None,
                ),
            ):
                readiness_contribution = weights["target_focus_capability_readiness"] * ready
                activity_contribution = weights["rarity_weighted_activity_similarity"] * work_similarity
                subfamily_contribution = weights["same_subfamily"] * float(same_subfamily)
                classification_contribution = weights["classification_proximity"] * grade_proximity
                score = (
                    readiness_contribution
                    + activity_contribution
                    + subfamily_contribution
                    + classification_contribution
                )
                record = {
                    "role_id": left["role_id"],
                    "pd_id": left["source_db_id"],
                    "pd_title": left["title"],
                    "neighbour_role_id": right["role_id"],
                    "neighbour_pd_id": right["source_db_id"],
                    "neighbour_title": right["title"],
                    "grade_step_dg": delta,
                    "readiness_rdy": round(ready, 12),
                    "work_similarity_cr": round(work_similarity, 12),
                    "same_subfamily_ssub": same_subfamily,
                    "grade_proximity_gp": round(grade_proximity, 12),
                    "readiness_contribution": round(readiness_contribution, 12),
                    "activity_contribution": round(activity_contribution, 12),
                    "subfamily_contribution": round(subfamily_contribution, 12),
                    "classification_contribution": round(classification_contribution, 12),
                    "score": round(score, 12),
                }
                heap = heaps[left["role_id"]]
                item = (score, right["role_id"], record)
                if len(heap) < neighbours_per_role:
                    heapq.heappush(heap, item)
                elif item[:2] > heap[0][:2]:
                    heapq.heapreplace(heap, item)
    output = []
    for role in roles:
        for neighbour_rank, item in enumerate(
            sorted(heaps[role["role_id"]], reverse=True), start=1
        ):
            item[2]["neighbour_rank"] = neighbour_rank
            output.append(item[2])
    neighbour_counts = Counter(row["role_id"] for row in output)
    scores = [float(row["score"]) for row in output]
    metrics = {
        "roles": len(roles),
        "edges": len(output),
        "neighbour_count_min": min(neighbour_counts.values()) if neighbour_counts else 0,
        "neighbour_count_max": max(neighbour_counts.values()) if neighbour_counts else 0,
        "score_min": min(scores) if scores else 0,
        "score_mean": sum(scores) / len(scores) if scores else 0,
        "score_max": max(scores) if scores else 0,
        "note": "Provisional transparent graph for data-control reconciliation; pathway-selection algorithm is not final.",
    }
    return output, metrics


def _generate_application_json(
    paths: Paths,
    roles: list[dict[str, str]],
    release: dict[str, Any],
    hierarchy: dict[str, dict[str, str]],
    algorithm_version: str,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    classification_by_raw, _ = _classification_data(paths)
    mappings: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in _read_csv(paths.canonical / "role_job_family_mappings.csv"):
        if row["role_id"]:
            mappings[row["role_id"]].append(row)
    caps: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in _read_csv(paths.canonical / "role_capabilities.csv"):
        caps[row["role_id"]].append(row)
    activities: dict[str, list[str]] = defaultdict(list)
    for row in _read_csv(paths.canonical / "role_activities.csv"):
        activities[row["role_id"]].append(row["activity_id"])
    requirements: dict[str, list[str]] = defaultdict(list)
    for row in _read_csv(paths.canonical / "essential_requirements.csv"):
        requirements[row["role_id"]].append(row["item_text"])
    purpose_by_db_id: dict[int, str] = {}
    connection = _connect_database(paths.database)
    for row in connection.execute(
        "SELECT position_description_id, section_text FROM pd_sections "
        "WHERE section_name = 'primary_purpose'"
    ):
        purpose_by_db_id[int(row["position_description_id"])] = row["section_text"]
    accountabilities: dict[int, list[str]] = defaultdict(list)
    for row in connection.execute(
        "SELECT position_description_id, item_text FROM pd_list_items "
        "WHERE section_name = 'key_accountabilities' "
        "ORDER BY position_description_id, sequence, id"
    ):
        accountabilities[int(row["position_description_id"])] .append(row["item_text"])
    connection.close()
    output = []
    ui_roles = []
    for role in roles:
        role_id = role["role_id"]
        db_id = int(role["source_db_id"])
        role_caps = caps[role_id]
        family = hierarchy.get(role_id, {})
        reference = classification_by_raw.get(role["classification_raw"])
        mapping_values = [
            {
                "rank": int(item["mapping_rank"]),
                "code": item["mapping_code"] or None,
                "name": item["mapping_name"] or None,
                "level": item["mapping_level"] or None,
                "validated": item["mapping_validated"] or None,
                "validation_notes": item["validation_notes"] or None,
                "framework_code_valid": item["framework_code_valid"] == "1",
            }
            for item in sorted(mappings[role_id], key=lambda item: int(item["mapping_rank"]))
        ]
        capability_values = [
            {
                "type": item["capability_type"],
                "framework": item["framework_name"],
                "group": item["capability_group"] or None,
                "capability": item["capability_name"],
                "code": item["capability_code"] or None,
                "level": item["required_level"],
                "normalized_level": float(item["normalized_level"]) if item["normalized_level"] else None,
                "description": item["description"] or None,
                "source_text": item["source_text"] or None,
            }
            for item in role_caps
        ]
        record = {
            "role_id": role_id,
            "source_release": release["source_release"],
            "build_version": BUILD_VERSION,
            "algorithm_version": algorithm_version,
            "source_db_id": db_id,
            "pd_id": role["pd_number"] or None,
            "title": role["title"],
            "classification": role["classification_raw"] or None,
            "classification_reference": reference or None,
            "job_family": family.get("job_family") or None,
            "sub_family": family.get("sub_family") or None,
            "specialisation": family.get("specialisation") or None,
            "job_family_mappings": mapping_values,
            "position_purpose": purpose_by_db_id.get(db_id),
            "key_accountabilities": accountabilities[db_id],
            "focus_capabilities": [
                item for item in capability_values if item["type"].lower() == "focus"
            ],
            "all_capabilities": capability_values,
            "activity_ids": activities[role_id],
            "essential_requirements": requirements[role_id],
        }
        output.append(record)
        ui_roles.append(
            {
                "id": role_id,
                "role_id": role_id,
                "position_description_no": role["pd_number"],
                "role_title": role["title"],
                "classification_display": reference.get("display_label") if reference else role["classification_raw"],
                "classification_abbreviation": reference.get("abbreviation") if reference else "",
                "classification_order": float(reference["seniority_order"]) if reference and reference.get("seniority_order") else None,
                "job_family_name": family.get("job_family"),
                "sub_family_name": family.get("sub_family"),
                "specialisation_name": family.get("specialisation"),
                "purpose": purpose_by_db_id.get(db_id),
                "capabilities": [item["capability"] for item in record["focus_capabilities"]],
                "activity_ids": activities[role_id],
                "source_filename": role["source_filename"],
            }
        )
    ui_payload = {
        "release": {
            "source_release": release["source_release"],
            "build_version": BUILD_VERSION,
            "algorithm_version": algorithm_version,
        },
        "summary": {
            "role_count": len(ui_roles),
            "job_family_count": len({row["job_family_name"] for row in ui_roles if row["job_family_name"]}),
            "grade_count": len({row["classification_display"] for row in ui_roles if row["classification_display"]}),
        },
        "roles": ui_roles,
    }
    return output, ui_payload


def _validate_and_report(
    paths: Paths,
    roles: list[dict[str, str]],
    role_by_id: dict[str, dict[str, str]],
    crosswalk_exceptions: list[dict[str, Any]],
    graph_metrics: dict[str, Any],
    release: dict[str, Any],
) -> dict[str, Any]:
    exceptions = list(crosswalk_exceptions)
    role_ids = [row["role_id"] for row in roles]
    for role_id, count in Counter(role_ids).items():
        if not role_id or count > 1:
            exceptions.append(
                {
                    "issue_type": "stable_role_id",
                    "severity": "error",
                    "role_id": role_id,
                    "record_key": role_id,
                    "description": "role_id is blank or duplicated",
                    "required_action": "assign one permanent unique role_id",
                }
            )
    pd_groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for role in roles:
        if not role["pd_number"]:
            exceptions.append(
                {
                    "issue_type": "missing_pd_number",
                    "severity": "warning",
                    "role_id": role["role_id"],
                    "record_key": role["source_db_id"],
                    "description": f"PD number missing for {role['title']}",
                    "required_action": "business owner supplies or confirms no PD number",
                }
            )
        else:
            pd_groups[role["pd_number"]].append(role)
    for pd_number, grouped in pd_groups.items():
        if len(grouped) > 1:
            exceptions.append(
                {
                    "issue_type": "duplicate_pd_number",
                    "severity": "error" if len({item["title"] for item in grouped}) > 1 else "warning",
                    "role_id": "|".join(item["role_id"] for item in grouped),
                    "record_key": pd_number,
                    "description": f"PD number used by {len(grouped)} role records: " + "; ".join(item["title"] for item in grouped),
                    "required_action": "business owner confirms collision, duplicate, or version relationship",
                }
            )
    applicability = _read_csv(paths.canonical / "capability_applicability.csv")
    for row in applicability:
        if row["applicability_status"] == "review_required":
            role = role_by_id[row["role_id"]]
            exceptions.append(
                {
                    "issue_type": "capability_applicability",
                    "severity": "error",
                    "role_id": row["role_id"],
                    "record_key": role["source_db_id"],
                    "description": f"No capability records for {role['title']}; missing versus not applicable is unresolved",
                    "required_action": "set explicit applicability_status to not_applicable or add capability records",
                }
            )
    for rule in _read_csv(paths.reference / "capability_level_rules.csv"):
        if rule["interpretation_status"] == "review_required":
            exceptions.append(
                {
                    "issue_type": "unknown_capability_level",
                    "severity": "warning",
                    "role_id": "",
                    "record_key": f"{rule['framework_name']}|{rule['raw_level']}",
                    "description": f"No numeric interpretation for raw level used {rule['assignment_count']} time(s)",
                    "required_action": "define a framework-specific normalisation rule or retain as non-numeric",
                }
            )
    for generated_name in (
        "pd_activity_rarity.csv",
        "pd_capability_profile_ac.csv",
        "pd_neighbour_graph.csv",
        "role_id_crosswalk.csv",
    ):
        generated_rows = _read_csv(paths.generated / generated_name)
        releases = {row.get("source_release", "") for row in generated_rows}
        versions = {row.get("build_version", "") for row in generated_rows}
        if releases != {release["source_release"]} or versions != {BUILD_VERSION}:
            exceptions.append(
                {
                    "issue_type": "stale_generated_file",
                    "severity": "error",
                    "role_id": "",
                    "record_key": generated_name,
                    "description": f"generated metadata does not match source release {release['source_release']} and build {BUILD_VERSION}",
                    "required_action": "run a complete rebuild and do not combine outputs from different releases",
                }
            )
    nodes = _read_csv(paths.reference / "job_family_nodes.csv")
    node_codes = {row["code"] for row in nodes}
    for node in nodes:
        if node["parent_code"] and node["parent_code"] not in node_codes:
            exceptions.append(
                {
                    "issue_type": "job_family_parent",
                    "severity": "error",
                    "role_id": "",
                    "record_key": node["code"],
                    "description": f"Unknown parent code {node['parent_code']}",
                    "required_action": "correct hierarchy parent reference",
                }
            )
    valid_classifications = {
        row["raw_label"] for row in _read_csv(paths.reference / "classification_aliases.csv")
    }
    for role in roles:
        if role["classification_raw"] and role["classification_raw"] not in valid_classifications:
            exceptions.append(
                {
                    "issue_type": "invalid_classification_reference",
                    "severity": "error",
                    "role_id": role["role_id"],
                    "record_key": role["classification_raw"],
                    "description": "Role classification has no alias/reference row",
                    "required_action": "add classification alias or correct role value",
                }
            )
    _write_csv(
        paths.qa / "exceptions.csv",
        [
            "issue_type",
            "severity",
            "role_id",
            "record_key",
            "description",
            "required_action",
        ],
        exceptions,
    )
    _write_csv(
        paths.qa / "business_owner_review.csv",
        [
            "issue_type",
            "severity",
            "role_id",
            "record_key",
            "description",
            "required_action",
        ],
        [row for row in exceptions if row["severity"] in {"error", "warning"}],
    )
    severity_counts = Counter(row["severity"] for row in exceptions)
    issue_counts = Counter(row["issue_type"] for row in exceptions)
    validation = {
        "source_release": release["source_release"],
        "build_version": BUILD_VERSION,
        "status": "review_required" if severity_counts["error"] else "passed_with_warnings" if severity_counts["warning"] else "passed",
        "role_count": len(roles),
        "exception_count": len(exceptions),
        "exceptions_by_severity": dict(severity_counts),
        "exceptions_by_type": dict(issue_counts),
        "graph_metrics": graph_metrics,
    }
    _write_json(paths.qa / "validation_summary.json", validation)
    _write_json(paths.qa / "graph_metrics.json", graph_metrics)
    legacy_counts = {}
    for name in ("pd_activity_rarity.csv", "pd_capability_profile_ac.csv", "pd_neighbour_graph.csv"):
        path = paths.legacy_dir / name
        if path.exists():
            rows = _read_csv(path)
            legacy_counts[name] = {
                "rows": len(rows),
                "roles": len({row["pd_id"] for row in rows}),
            }
    reconciliation = {
        "source_release": release["source_release"],
        "catalogue": {"roles": len(roles)},
        "activity_workbook": {
            "roles": len({row["role_id"] for row in _read_csv(paths.canonical / "role_activities.csv")}),
            "relationship_rows": len(_read_csv(paths.canonical / "role_activities.csv")),
        },
        "legacy_generated_files": legacy_counts,
        "generated": {
            "capability_profile_roles": len({row["role_id"] for row in _read_csv(paths.generated / "pd_capability_profile_ac.csv")}),
            "activity_rarity_roles": len({row["role_id"] for row in _read_csv(paths.generated / "pd_activity_rarity.csv")}),
            "neighbour_graph_roles": len({row["role_id"] for row in _read_csv(paths.generated / "pd_neighbour_graph.csv")}),
        },
    }
    _write_json(paths.qa / "reconciliation.json", reconciliation)
    return validation


def _write_sqlite(
    paths: Paths,
    release: dict[str, Any],
    config: dict[str, Any],
    application_rows: list[dict[str, Any]],
) -> None:
    target = paths.generated / "career_pathways.sqlite3"
    temporary = target.with_suffix(".tmp.sqlite3")
    if temporary.exists():
        temporary.unlink()
    connection = sqlite3.connect(temporary)
    try:
        table_sources = {
            "roles": paths.canonical / "roles.csv",
            "classification_aliases": paths.reference / "classification_aliases.csv",
            "job_family_nodes": paths.reference / "job_family_nodes.csv",
            "role_job_family_mappings": paths.canonical / "role_job_family_mappings.csv",
            "capability_frameworks": paths.reference / "capability_frameworks.csv",
            "capabilities": paths.reference / "capabilities.csv",
            "capability_level_rules": paths.reference / "capability_level_rules.csv",
            "role_capabilities": paths.canonical / "role_capabilities.csv",
            "capability_applicability": paths.canonical / "capability_applicability.csv",
            "activities": paths.reference / "activities.csv",
            "role_activities": paths.canonical / "role_activities.csv",
            "essential_requirements": paths.canonical / "essential_requirements.csv",
            "adjacency_rules": paths.reference / "adjacency_rules.csv",
            "activity_rarity": paths.generated / "pd_activity_rarity.csv",
            "role_capability_profiles": paths.generated / "pd_capability_profile_ac.csv",
            "role_neighbours": paths.generated / "pd_neighbour_graph.csv",
            "role_id_crosswalk": paths.generated / "role_id_crosswalk.csv",
        }
        for table, path in table_sources.items():
            rows = _read_csv(path)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                fields = csv.DictReader(handle).fieldnames or []
            columns = ", ".join(f'"{field}" TEXT' for field in fields)
            connection.execute(f'CREATE TABLE "{table}" ({columns})')
            if rows:
                placeholders = ", ".join("?" for _ in fields)
                connection.executemany(
                    f'INSERT INTO "{table}" VALUES ({placeholders})',
                    [[row.get(field, "") for field in fields] for row in rows],
                )
        connection.execute("CREATE TABLE build_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        connection.executemany(
            "INSERT INTO build_metadata(key, value) VALUES (?, ?)",
            [
                ("source_release", release["source_release"]),
                ("build_version", BUILD_VERSION),
                ("algorithm_version", str(config["algorithm_version"])),
                ("algorithm_config_json", json.dumps(config, sort_keys=True)),
                (
                    "snapshot_scope",
                    "complete structured inputs, derived features, scoring configuration, and recommendation outputs",
                ),
            ],
        )

        connection.execute(
            "CREATE TABLE algorithm_parameters ("
            "parameter_key TEXT PRIMARY KEY, value TEXT NOT NULL, value_type TEXT NOT NULL)"
        )

        def parameter_rows(value: Any, prefix: str = "") -> list[tuple[str, str, str]]:
            rows: list[tuple[str, str, str]] = []
            if isinstance(value, dict):
                for key in sorted(value):
                    nested = f"{prefix}.{key}" if prefix else key
                    rows.extend(parameter_rows(value[key], nested))
            else:
                value_type = "boolean" if isinstance(value, bool) else type(value).__name__
                stored = json.dumps(value) if isinstance(value, (bool, list)) else str(value)
                rows.append((prefix, stored, value_type))
            return rows

        connection.executemany(
            "INSERT INTO algorithm_parameters(parameter_key, value, value_type) VALUES (?, ?, ?)",
            parameter_rows(config),
        )

        connection.execute(
            "CREATE TABLE activity_statistics AS "
            "SELECT activity_id, activity_label, roles_with_activity, total_active_roles, "
            "share, rarity_weight, source_release, build_version, algorithm_version "
            "FROM activity_rarity GROUP BY activity_id, activity_label, roles_with_activity, "
            "total_active_roles, share, rarity_weight, source_release, build_version, algorithm_version"
        )
        connection.execute(
            "CREATE UNIQUE INDEX activity_statistics_activity_uq "
            "ON activity_statistics(activity_id)"
        )

        connection.execute(
            "CREATE TABLE release_files ("
            "file_class TEXT NOT NULL, path TEXT NOT NULL, sha256 TEXT NOT NULL, "
            "PRIMARY KEY(file_class, path))"
        )
        release_files = [
            (file_class, row["path"], row["sha256"])
            for file_class, key in (("upstream_source", "source_files"), ("canonical_input", "canonical_files"))
            for row in release.get(key, [])
        ]
        connection.executemany(
            "INSERT INTO release_files(file_class, path, sha256) VALUES (?, ?, ?)",
            release_files,
        )

        connection.execute(
            "CREATE TABLE role_application_payloads ("
            "role_id TEXT PRIMARY KEY, payload_json TEXT NOT NULL)"
        )
        connection.executemany(
            "INSERT INTO role_application_payloads(role_id, payload_json) VALUES (?, ?)",
            [
                (row["role_id"], json.dumps(row, ensure_ascii=False, sort_keys=True))
                for row in application_rows
            ],
        )
        connection.execute(
            "CREATE TABLE role_purposes ("
            "role_id TEXT PRIMARY KEY, position_purpose TEXT NOT NULL)"
        )
        connection.executemany(
            "INSERT INTO role_purposes(role_id, position_purpose) VALUES (?, ?)",
            [
                (row["role_id"], row["position_purpose"])
                for row in application_rows
                if row.get("position_purpose")
            ],
        )
        connection.execute(
            "CREATE TABLE key_accountabilities ("
            "role_id TEXT NOT NULL, sequence INTEGER NOT NULL, item_text TEXT NOT NULL, "
            "PRIMARY KEY(role_id, sequence))"
        )
        connection.executemany(
            "INSERT INTO key_accountabilities(role_id, sequence, item_text) VALUES (?, ?, ?)",
            [
                (row["role_id"], sequence, item)
                for row in application_rows
                for sequence, item in enumerate(row.get("key_accountabilities", []), start=1)
            ],
        )

        connection.execute("CREATE UNIQUE INDEX roles_role_id_uq ON roles(role_id)")
        connection.execute("CREATE INDEX role_capabilities_role_idx ON role_capabilities(role_id)")
        connection.execute("CREATE INDEX role_activities_role_idx ON role_activities(role_id)")
        connection.execute("CREATE INDEX activity_rarity_role_idx ON activity_rarity(role_id)")
        connection.execute("CREATE INDEX activity_rarity_activity_idx ON activity_rarity(activity_id)")
        connection.execute("CREATE INDEX role_profiles_role_idx ON role_capability_profiles(role_id)")
        connection.execute("CREATE INDEX role_neighbours_role_rank_idx ON role_neighbours(role_id, neighbour_rank)")
        connection.execute("CREATE INDEX role_neighbours_target_idx ON role_neighbours(neighbour_role_id)")
        connection.execute("CREATE INDEX key_accountabilities_role_idx ON key_accountabilities(role_id)")

        connection.execute(
            "CREATE VIEW role_activity_audit AS "
            "SELECT r.role_id, r.pd_number, r.title, ra.activity_rank, ra.activity_id, "
            "a.canonical_plain_label AS activity_name, ar.roles_with_activity, "
            "ar.total_active_roles, ar.share, ar.rarity_weight, "
            "ra.assignment_gaps, ra.retry_count, ra.fallback_used "
            "FROM roles r JOIN role_activities ra ON ra.role_id = r.role_id "
            "LEFT JOIN activities a ON a.activity_id = ra.activity_id "
            "LEFT JOIN activity_rarity ar ON ar.role_id = ra.role_id AND ar.activity_id = ra.activity_id"
        )
        connection.execute(
            "CREATE VIEW role_capability_audit AS "
            "SELECT r.role_id, r.pd_number, r.title, p.framework, p.capability_code, "
            "p.capability, p.raw_level, p.level AS normalized_level, p.focus, "
            "p.source_release, p.build_version, p.algorithm_version "
            "FROM roles r JOIN role_capability_profiles p ON p.role_id = r.role_id"
        )
        connection.execute(
            "CREATE VIEW role_neighbour_audit AS "
            "SELECT n.*, source.pd_number AS source_pd_number, "
            "target.pd_number AS neighbour_pd_number_business "
            "FROM role_neighbours n "
            "JOIN roles source ON source.role_id = n.role_id "
            "JOIN roles target ON target.role_id = n.neighbour_role_id"
        )

        connection.execute(
            "CREATE TABLE dataset_inventory ("
            "dataset_name TEXT PRIMARY KEY, row_count INTEGER NOT NULL)"
        )
        dataset_names = list(table_sources) + [
            "algorithm_parameters",
            "activity_statistics",
            "release_files",
            "role_application_payloads",
            "role_purposes",
            "key_accountabilities",
        ]
        connection.executemany(
            "INSERT INTO dataset_inventory(dataset_name, row_count) VALUES (?, ?)",
            [
                (name, connection.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0])
                for name in dataset_names
            ],
        )
        connection.commit()
    finally:
        connection.close()
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary.replace(target)


def build_outputs(paths: Paths) -> dict[str, Any]:
    paths.generated.mkdir(parents=True, exist_ok=True)
    paths.qa.mkdir(parents=True, exist_ok=True)
    release = json.loads((paths.canonical / "source_release.json").read_text(encoding="utf-8"))
    config = json.loads((paths.canonical / "build_config.json").read_text(encoding="utf-8"))
    maintained_files = sorted(paths.canonical.glob("*.csv")) + sorted(paths.reference.glob("*.csv")) + [
        paths.canonical / "build_config.json"
    ]
    canonical_release_hash = hashlib.sha256(
        "|".join(f"{path.relative_to(paths.root)}:{_sha256(path)}" for path in maintained_files).encode()
    ).hexdigest()
    release["source_release"] = canonical_release_hash[:16]
    release["build_version"] = BUILD_VERSION
    release["algorithm_version"] = config["algorithm_version"]
    previous_manifest = None
    manifest_path = paths.generated / "manifest.json"
    if manifest_path.exists():
        previous_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    release["canonical_files"] = [
        {"path": str(path.relative_to(paths.root)), "sha256": _sha256(path)}
        for path in maintained_files
    ]
    _write_json(paths.canonical / "source_release.json", release)
    roles, role_by_id = _role_data(paths)
    if not roles:
        raise ValueError("roles.csv is empty")
    crosswalk, crosswalk_exceptions = _legacy_crosswalk(paths, roles)
    for row in crosswalk:
        row["source_release"] = release["source_release"]
        row["build_version"] = BUILD_VERSION
    _write_csv(
        paths.generated / "role_id_crosswalk.csv",
        [
            "role_id",
            "source_system",
            "legacy_identifier_type",
            "legacy_identifier",
            "match_status",
            "match_basis",
            "pd_number",
            "title",
            "classification",
            "source_release",
            "build_version",
        ],
        crosswalk,
    )
    rarity_rows, activity_shares = _generate_activity_rarity(paths, roles, role_by_id)
    for row in rarity_rows:
        row["source_release"] = release["source_release"]
        row["build_version"] = BUILD_VERSION
        row["algorithm_version"] = config["algorithm_version"]
    _write_csv(
        paths.generated / "pd_activity_rarity.csv",
        [
            "role_id",
            "pd_id",
            "pd_title",
            "activity_id",
            "activity_label",
            "roles_with_activity",
            "total_active_roles",
            "share",
            "rarity_weight",
            "source_release",
            "build_version",
            "algorithm_version",
        ],
        rarity_rows,
    )
    profile_rows, focus = _generate_capability_profile(paths, role_by_id)
    for row in profile_rows:
        row["source_release"] = release["source_release"]
        row["build_version"] = BUILD_VERSION
        row["algorithm_version"] = config["algorithm_version"]
    _write_csv(
        paths.generated / "pd_capability_profile_ac.csv",
        [
            "role_id",
            "pd_id",
            "pd_title",
            "pd_grade",
            "framework",
            "capability_code",
            "capability",
            "raw_level",
            "level",
            "focus",
            "source_release",
            "build_version",
            "algorithm_version",
        ],
        profile_rows,
    )
    classification_by_raw, classification_ranks = _classification_data(paths)
    _, hierarchy = _primary_family_maps(paths)
    graph_rows, graph_metrics = _generate_neighbours(
        paths,
        roles,
        role_by_id,
        focus,
        activity_shares,
        hierarchy,
        classification_by_raw,
        classification_ranks,
        int(config["neighbours_per_role"]),
        {key: float(value) for key, value in config["weights"].items()},
    )
    for row in graph_rows:
        row["source_release"] = release["source_release"]
        row["build_version"] = BUILD_VERSION
        row["algorithm_version"] = config["algorithm_version"]
    _write_csv(
        paths.generated / "pd_neighbour_graph.csv",
        [
            "role_id",
            "pd_id",
            "pd_title",
            "neighbour_role_id",
            "neighbour_pd_id",
            "neighbour_title",
            "neighbour_rank",
            "grade_step_dg",
            "readiness_rdy",
            "work_similarity_cr",
            "same_subfamily_ssub",
            "grade_proximity_gp",
            "readiness_contribution",
            "activity_contribution",
            "subfamily_contribution",
            "classification_contribution",
            "score",
            "source_release",
            "build_version",
            "algorithm_version",
        ],
        graph_rows,
    )
    application_rows, ui_payload = _generate_application_json(
        paths, roles, release, hierarchy, str(config["algorithm_version"])
    )
    _write_json(
        paths.generated
        / "pd-data-for-claude-career-pathways-with-classification-order.json",
        application_rows,
    )
    _write_json(paths.generated / "career-explorer-payload.json", ui_payload)
    _write_sqlite(paths, release, config, application_rows)
    validation = _validate_and_report(
        paths,
        roles,
        role_by_id,
        crosswalk_exceptions,
        graph_metrics,
        release,
    )
    if previous_manifest:
        previous_outputs = {
            Path(row["path"]).name: row["sha256"]
            for row in previous_manifest.get("outputs", [])
        }
        current_generated_files = [
            path for path in sorted(paths.generated.iterdir())
            if path.is_file() and path.name != "manifest.json"
        ]
        changed = [
            path.name
            for path in current_generated_files
            if previous_outputs.get(path.name) != _sha256(path)
        ]
        removed = sorted(set(previous_outputs) - {path.name for path in current_generated_files})
        comparison = {
            "previous_source_release": previous_manifest.get("source_release"),
            "current_source_release": release["source_release"],
            "changed_outputs": changed,
            "removed_outputs": removed,
            "graph_edge_change": graph_metrics["edges"]
            - int(previous_manifest.get("graph_metrics", {}).get("edges", graph_metrics["edges"])),
        }
    else:
        comparison = {
            "previous_source_release": None,
            "current_source_release": release["source_release"],
            "changed_outputs": [
                path.name
                for path in sorted(paths.generated.iterdir())
                if path.is_file() and path.name != "manifest.json"
            ],
            "removed_outputs": [],
            "graph_edge_change": None,
        }
    _write_json(paths.qa / "release_comparison.json", comparison)
    generated_manifest = {
        "source_release": release["source_release"],
        "build_version": BUILD_VERSION,
        "algorithm_version": config["algorithm_version"],
        "built_at_utc": datetime.now(timezone.utc).isoformat(),
        "graph_metrics": graph_metrics,
        "outputs": [
            {
                "path": str(path.relative_to(paths.root)),
                "sha256": _sha256(path),
                "bytes": path.stat().st_size,
            }
            for path in sorted(paths.generated.iterdir())
            if path.is_file() and path.name != "manifest.json"
        ],
    }
    _write_json(paths.generated / "manifest.json", generated_manifest)
    return {"release": release, "validation": validation, "manifest": generated_manifest}


def _write_docs(paths: Paths, result: dict[str, Any]) -> None:
    validation = result["validation"]
    release = result["release"]
    inventory = """# Career Pathways data inventory

| File / dataset | Purpose | Rows / shape | Identifier | Classification | Decision |
|---|---|---:|---|---|---|
| `pd-management-bulk-fixed.sqlite3` / `position_descriptions` | Current role catalogue and extracted PD facts | 1,171 roles | database `id`; PD number is non-unique | Maintained upstream | Import into canonical source tables; retain provenance |
| `activity-assigned-clustered-role-links-revised-high-all.xlsx` / Summary | Workbook build summary | 8 × 2 | n/a | Generated summary | Retain for audit |
| same workbook / Role Activity Links | Role-to-activity relationships | 13,936 data rows × 14 | database `PD ID` + `Activity ID` | Maintained/curated relationship output | Import into `role_activities.csv` |
| same workbook / Role Summary | Per-role activity roll-up | 1,168 data rows × 13 | database `PD ID` | Duplicated summary | Regenerate/retire as authority |
| same workbook / Vocabulary Usage | Activity vocabulary statistics | 1,285 data rows × 15 | `Activity ID` | Mixed reference/calculated | Split activity definitions from generated usage measures |
| same workbook / Role Activity Matrix | Wide role/activity matrix | 1,168 data rows × 1,288 | database `PD ID` | Generated duplicate | Regenerate/retire as authority |
| `classification-references-20260716-145030.csv/json` | Classification aliases and ordering | 69 records | raw classification label | Maintained reference, duplicated formats | Canonical CSV in `data/reference`; JSON retired |
| combined career-pathways JSON | Application bundle of role facts | 1,171 roles | PD number/title only; neither unique | Generated from SQLite | Regenerate with `role_id` and release metadata |
| `pd_activity_rarity.csv` | Per-role activity rarity | 13,001 rows; 1,087 roles | database-like `pd_id` | Calculated, stale/incomplete | Regenerate from canonical role activities |
| `pd_capability_profile_ac.csv` | Flattened numeric capability profile | 19,878 rows; 1,099 roles | positional `pd_id` 0–1,098 | Calculated from older catalogue | Regenerate with stable `role_id`, raw and interpreted levels |
| `pd_neighbour_graph.csv` | Candidate role adjacency | 42,582 rows; 1,099 roles | positional `pd_id` 0–1,098 | Calculated from older catalogue | Regenerate provisionally; final pathway algorithm remains out of scope |

The OneDrive originals remain unmodified. The CSVs under `data/source/onedrive_exports/` are explicitly labelled browser exports and are used only for reconciliation, not as canonical sources.

`walk-algorithm-brief.md` was referenced by the supplied brief but was not present in the OneDrive folder or repository. Its absence is recorded rather than silently substituting an assumed algorithm definition.
"""
    (paths.docs / "career-pathways-data-inventory.md").write_text(inventory, encoding="utf-8")
    _write_csv(
        paths.qa / "source_inventory.csv",
        [
            "file",
            "dataset",
            "purpose",
            "rows",
            "columns",
            "fields",
            "apparent_primary_key",
            "foreign_keys_dependencies",
            "catalogue_date_or_version",
            "classification",
            "data_quality_problems",
            "decision",
        ],
        [
            {
                "file": "output/pd-management-bulk-fixed.sqlite3",
                "dataset": "position_descriptions and normalized child tables",
                "purpose": "current catalogue and extracted PD facts",
                "rows": 1171,
                "columns": "normalized schema",
                "fields": "id; source_filename; role_title; position_description_no; classification_grade_band; provenance and child facts",
                "apparent_primary_key": "position_descriptions.id",
                "foreign_keys_dependencies": "all PD child tables reference position_description_id",
                "catalogue_date_or_version": "current local catalogue used for 2026-08-12 exports",
                "classification": "maintained upstream",
                "data_quality_problems": "12 missing PD numbers; five duplicated PD numbers; titles non-unique",
                "decision": "import into separate canonical CSV source tables",
            },
            {
                "file": "activity-assigned-clustered-role-links-revised-high-all.xlsx",
                "dataset": "Summary",
                "purpose": "workbook metrics",
                "rows": 8,
                "columns": 2,
                "fields": "Metric; Value",
                "apparent_primary_key": "Metric",
                "foreign_keys_dependencies": "summarises workbook sheets",
                "catalogue_date_or_version": "modified 2026-08-12",
                "classification": "generated summary",
                "data_quality_problems": "not authoritative",
                "decision": "retain for provenance",
            },
            {
                "file": "activity-assigned-clustered-role-links-revised-high-all.xlsx",
                "dataset": "Role Activity Links",
                "purpose": "role-to-activity assignments",
                "rows": 13936,
                "columns": 14,
                "fields": "PD ID; PD Number; Title; Activity Rank; Activity ID; Canonical Match Label; Canonical Plain Label; Discriminating; Cluster ID; Raw Count; Cluster Member Phrases; Assignment Gaps; Retry Count; Fallback Used",
                "apparent_primary_key": "PD ID + Activity ID",
                "foreign_keys_dependencies": "PD ID -> database role; Activity ID -> vocabulary",
                "catalogue_date_or_version": "1,168-role release; modified 2026-08-12",
                "classification": "curated relationship source with calculated diagnostics",
                "data_quality_problems": "three catalogue role IDs absent; diagnostics mixed with facts",
                "decision": "split activities and role_activities",
            },
            {
                "file": "activity-assigned-clustered-role-links-revised-high-all.xlsx",
                "dataset": "Role Summary",
                "purpose": "role activity roll-up",
                "rows": 1168,
                "columns": 13,
                "fields": "PD ID; PD Number; Title; source/excluded/extracted/assigned counts; retries; fallback; activity arrays; gaps",
                "apparent_primary_key": "PD ID",
                "foreign_keys_dependencies": "derived from Role Activity Links",
                "catalogue_date_or_version": "1,168-role release",
                "classification": "calculated duplicate",
                "data_quality_problems": "not same role population as 1,171-role catalogue",
                "decision": "regenerate; retire as authority",
            },
            {
                "file": "activity-assigned-clustered-role-links-revised-high-all.xlsx",
                "dataset": "Vocabulary Usage",
                "purpose": "activity vocabulary and usage statistics",
                "rows": 1285,
                "columns": 15,
                "fields": "Activity ID; canonical labels; discriminating; cluster fields; occurrence and assignment measures",
                "apparent_primary_key": "Activity ID",
                "foreign_keys_dependencies": "referenced by Role Activity Links and Matrix",
                "catalogue_date_or_version": "modified 2026-08-12",
                "classification": "mixed reference and calculated",
                "data_quality_problems": "definition and usage measures combined",
                "decision": "retain definitions; regenerate usage measures",
            },
            {
                "file": "activity-assigned-clustered-role-links-revised-high-all.xlsx",
                "dataset": "Role Activity Matrix",
                "purpose": "wide application/analysis representation",
                "rows": 1168,
                "columns": 1288,
                "fields": "PD ID; PD Number; Title; one column per activity",
                "apparent_primary_key": "PD ID",
                "foreign_keys_dependencies": "derived from role_activities and activities",
                "catalogue_date_or_version": "1,168-role release",
                "classification": "generated duplicate",
                "data_quality_problems": "wide format and stale population make maintenance unsafe",
                "decision": "regenerate only when needed; retire as source",
            },
            {
                "file": "classification-references-20260716-145030.csv/json",
                "dataset": "classification references",
                "purpose": "classification aliases, cohorts and seniority order",
                "rows": 69,
                "columns": 10,
                "fields": "raw_label; usage_count; display_label; abbreviation; cohort; seniority_order; enterprise_agreement; active; notes; updated_at",
                "apparent_primary_key": "raw_label",
                "foreign_keys_dependencies": "roles.classification_raw -> raw_label",
                "catalogue_date_or_version": "2026-07-16",
                "classification": "maintained reference duplicated in two formats",
                "data_quality_problems": "JSON and CSV duplicate the database table",
                "decision": "canonical reference CSV; JSON retired",
            },
            {
                "file": "pd-data-for-claude-career-pathways-with-classification-order.json",
                "dataset": "combined roles",
                "purpose": "application/research bundle",
                "rows": 1171,
                "columns": 12,
                "fields": "pd_id; title; classification; classification_reference; flattened and ranked family mapping; purpose; accountabilities; focus/all capabilities",
                "apparent_primary_key": "none",
                "foreign_keys_dependencies": "generated from SQLite normalized tables",
                "catalogue_date_or_version": "modified 2026-08-12",
                "classification": "generated application bundle",
                "data_quality_problems": "PD numbers missing/duplicate; title non-unique; flattened family fields keep only first mapping",
                "decision": "regenerate with role_id and release metadata",
            },
            {
                "file": "pd_activity_rarity.csv",
                "dataset": "activity rarity",
                "purpose": "calculated role/activity weights",
                "rows": 13001,
                "columns": 6,
                "fields": "pd_id; pd_title; activity_id; activity_label; share; rarity_weight",
                "apparent_primary_key": "pd_id + activity_id",
                "foreign_keys_dependencies": "role catalogue and role activities",
                "catalogue_date_or_version": "1,087-role subset; modified 2026-08-12",
                "classification": "calculated",
                "data_quality_problems": "does not cover current catalogue; no release metadata",
                "decision": "regenerate",
            },
            {
                "file": "pd_capability_profile_ac.csv",
                "dataset": "capability profile",
                "purpose": "flattened numeric capability input",
                "rows": 19878,
                "columns": 5,
                "fields": "pd_id; pd_title; pd_grade; capability; level",
                "apparent_primary_key": "positional pd_id + capability",
                "foreign_keys_dependencies": "older role catalogue and capability assignments",
                "catalogue_date_or_version": "1,099-role older release",
                "classification": "calculated",
                "data_quality_problems": "positional IDs; raw framework/code/level lost; stale population",
                "decision": "regenerate with stable IDs and retained source values",
            },
            {
                "file": "pd_neighbour_graph.csv",
                "dataset": "neighbour graph",
                "purpose": "role adjacency candidates",
                "rows": 42582,
                "columns": 8,
                "fields": "pd_id; pd_title; neighbour_pd_id; neighbour_title; grade_step_dg; readiness_rdy; work_similarity_cr; same_subfamily_ssub",
                "apparent_primary_key": "positional pd_id + neighbour_pd_id",
                "foreign_keys_dependencies": "older capability profile, activities and family mappings",
                "catalogue_date_or_version": "1,099-role older release",
                "classification": "calculated",
                "data_quality_problems": "positional IDs and no reproducible release metadata",
                "decision": "regenerate provisionally; final algorithm deferred",
            },
            {
                "file": "walk-algorithm-brief.md",
                "dataset": "pathway mechanism overview",
                "purpose": "algorithm documentation referenced by brief",
                "rows": "missing",
                "columns": "n/a",
                "fields": "n/a",
                "apparent_primary_key": "n/a",
                "foreign_keys_dependencies": "would inform generated graph/pathway rules",
                "catalogue_date_or_version": "unknown",
                "classification": "referenced but not supplied",
                "data_quality_problems": "not present in repository or OneDrive folder",
                "decision": "record as missing; do not infer final algorithm",
            },
        ],
    )
    model = """# Career Pathways canonical data model

The logical model is relational. Maintained role facts and relationships live in `data/source/canonical/`; taxonomies and interpretation rules live in `data/reference/`; all algorithm and application artifacts live in `data/generated/`.

`roles(role_id)` is the parent of role capabilities, activities, requirements, applicability, family mappings and generated graph edges. `role_id` is a stored UUID. The initial UUID is deterministically bootstrapped from the catalogue database primary key, then persisted in `roles.csv`; later imports reuse the persisted mapping.

Core entities:

- `roles`: stable identity, PD business number, title, classification alias and provenance.
- `classification_aliases`: raw-to-display classification mapping, cohort and seniority order.
- `job_family_nodes`: hierarchy nodes with parent codes; `role_job_family_mappings` retains ranks and alternatives.
- `capability_frameworks`, `capabilities`, `role_capabilities`: original framework, code, name and raw level. `normalized_level` is a separate calculated interpretation.
- `capability_level_rules`: observed raw formats, numeric interpretations and review status.
- `capability_applicability`: explicit `applicable`, `not_applicable` or `review_required`; absence is never silently treated as not applicable.
- `activities`, `role_activities`: vocabulary and role assignments.
- `essential_requirements`: mandatory/essential items extracted from supplied PDs.
- `adjacency_rules`: reserved maintained family/occupation rules; currently empty rather than guessed.
- `source_release.json`, `build_config.json`: input hashes, counts, build version and calculation configuration.
- `activity_rarity`: stored corpus share and rarity weight for every role/activity assignment.
- `role_capability_profiles`: the exact normalized capability features consumed by scoring.
- `role_neighbours`: ranked recommendations with every raw component, weighted contribution and final score.
- `algorithm_parameters`, `build_metadata`, `release_files`: the exact rules, versions and input hashes needed to reproduce a release.
- `role_application_payloads`, `role_purposes`, `key_accountabilities`: the complete structured application data included in the published snapshot.

The SQLite file is the complete release snapshot. Generated CSV and JSON files are compatibility and interchange exports, not the sole location of any algorithm input or result. Audit-friendly views (`role_activity_audit`, `role_capability_audit`, and `role_neighbour_audit`) expose readable role-level records without requiring joins.

Generated graph scoring is transparent and provisional. It exists to make all outputs reproducible from one release; final pathway-selection tuning is deliberately deferred.
"""
    (paths.docs / "career-pathways-data-model.md").write_text(model, encoding="utf-8")
    mapping = """# Existing-to-canonical field mapping

| Existing source | Existing field | Canonical target | Transformation / control |
|---|---|---|---|
| SQLite `position_descriptions` | `id` | `roles.source_db_id`; role-ID crosswalk | Bootstrap lineage only; never exposed as stable identity |
| SQLite `position_descriptions` | `position_description_no` | `roles.pd_number` | Retained business identifier; missing/duplicates reported |
| SQLite `position_descriptions` | title, classification, source fields | `roles.*` | Direct import with provenance |
| Classification references | raw/display/abbreviation/cohort/order | `classification_aliases.*` | CSV is canonical reference representation |
| Active job-family entries | code/name/level/parent | `job_family_nodes.*` | Direct relational hierarchy import |
| Active role mappings | rank/code/name/validation | `role_job_family_mappings.*` | All ranked alternatives retained; no flattening loss |
| PD capabilities | framework/name/code/type/raw level | `capabilities`, `role_capabilities` | Framework-specific source values retained |
| PD capabilities | raw level | `role_capabilities.normalized_level` | Separate calculated interpretation via level rules |
| Activity workbook `PD ID` | role link | `role_activities.role_id` | Join through database ID crosswalk |
| Activity workbook `Activity ID` and labels | activity definition/link | `activities`, `role_activities` | Vocabulary separated from assignments |
| Legacy capability/graph `pd_id` | 0–1,098 array position | `role_id_crosswalk` | Unique title+classification matches only; ambiguous rows reported |
| Legacy activity rarity `pd_id` | database ID | `role_id_crosswalk` | Requires database ID and title agreement |
| Combined JSON flattened family fields | first mapping only | generated application JSON | Rebuilt from ranked mappings; flattened primary fields remain compatibility views |
"""
    (paths.docs / "career-pathways-field-mapping.md").write_text(mapping, encoding="utf-8")
    maintenance = f"""# Maintaining Career Pathways data

## One-command rebuild

After editing the maintained CSV tables, run:

```powershell
python scripts/build_career_pathways.py build
```

To re-import the current catalogue database and activity workbook, preserving existing `role_id` values, then rebuild everything:

```powershell
python scripts/build_career_pathways.py all
```

## Add or change a role

1. Add or update one row in `data/source/canonical/roles.csv`. Assign a new UUID once for a genuinely new role; never reuse or renumber an existing `role_id`.
2. Update its classification, ranked job-family mappings, capability applicability, capabilities, activities and essential requirements in the corresponding canonical tables.
3. Retire a role by setting `record_status=retired`; do not delete or recycle its identity.
4. Run the build command.
5. Review `data/qa/validation_summary.json`, `exceptions.csv`, `business_owner_review.csv`, `reconciliation.json` and `change_report.md`.
6. Publish only when exceptions are accepted/resolved and all generated files share source release `{release['source_release']}`.

The build performs a complete rebuild. Generated CSV, JSON and SQLite artifacts must never be manually edited.

The SQLite database is a self-contained release snapshot containing canonical inputs, derived algorithm features, configuration, lineage hashes and recommendation outputs. Verify it with:

```powershell
python scripts/verify_career_pathways_snapshot.py
```
"""
    (paths.docs / "career-pathways-maintenance.md").write_text(maintenance, encoding="utf-8")
    change = f"""# Career Pathways build change report

- Source release: `{release['source_release']}`
- Build version: `{BUILD_VERSION}`
- Roles: {validation['role_count']:,}
- Validation status: `{validation['status']}`
- Exceptions: {validation['exception_count']:,}

Important decisions:

- Stable UUIDs are persisted and are not derived from array position, title or PD number.
- Database ID is used only to bootstrap and crosswalk the current catalogue.
- Ranked job-family alternatives are retained; flattened family values are generated compatibility fields.
- Capability raw levels and frameworks are preserved. Numeric values are interpretations, never replacements.
- Capability absence remains `review_required` until a business owner explicitly declares `not_applicable` or supplies data.
- Activity rarity, capability profiles, graph component scores, algorithm parameters and application payloads are stored in the SQLite release snapshot as well as exported where needed.
- Every neighbour score is the sum of four stored weighted contributions and has an explicit rank and algorithm version.
- The graph scoring configuration is provisional because final pathway selection is outside this brief.
"""
    (paths.qa / "change_report.md").write_text(change, encoding="utf-8")
    decision_log = """# Career Pathways decision log

1. **Identity:** use a stored UUID `role_id`. Database IDs are bootstrap lineage; array positions, titles and PD numbers are never technical keys.
2. **Authority:** canonical CSV tables are the maintained pathway source. The SQLite release is the complete, immutable publication snapshot; JSON and CSV outputs are compatibility exports.
3. **Mappings:** retain every ranked job-family mapping. Primary flattened fields are generated compatibility views only.
4. **Capabilities:** preserve framework, code, name and raw level. Numeric normalisation is separate and reviewable.
5. **Applicability:** capability absence is `review_required` unless a business owner explicitly records `not_applicable`.
6. **Ambiguity:** legacy positional IDs are crosswalked only on unambiguous evidence. No silent title-based collision resolution.
7. **Rebuild:** complete rebuilds are used at the present catalogue size.
8. **Algorithm scope:** the generated graph is a transparent provisional reconciliation artifact. Final pathway-selection tuning awaits the missing algorithm brief and business decisions.
9. **Auditability:** every structured algorithm input, parameter, derived feature, score contribution and recommendation is persisted in the versioned SQLite snapshot.
"""
    (paths.docs / "career-pathways-decision-log.md").write_text(decision_log, encoding="utf-8")


def build_all(paths: Paths) -> dict[str, Any]:
    import_summary = import_sources(paths)
    result = build_outputs(paths)
    result["import"] = import_summary
    _write_docs(paths, result)
    return result


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build canonical Career Pathways data")
    parser.add_argument("command", choices=("import", "build", "all"), nargs="?", default="all")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[3])
    parser.add_argument(
        "--database",
        type=Path,
        help="Operational database to import (defaults to the legacy fixed snapshot)",
    )
    args = parser.parse_args(argv)
    paths = Paths(
        args.root.resolve(),
        args.database.resolve() if args.database else None,
    )
    if args.command == "import":
        result: dict[str, Any] = import_sources(paths)
    elif args.command == "build":
        result = build_outputs(paths)
        _write_docs(paths, result)
    else:
        result = build_all(paths)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
