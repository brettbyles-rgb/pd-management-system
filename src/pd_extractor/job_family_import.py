from __future__ import annotations

import argparse
import json
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .database import connect_database, database_statistics, initialise_database


FRAMEWORK_SHEET = "job_family_power_query"
MAPPING_SHEET = "job_mapping_input"


@dataclass(frozen=True)
class JobFamilyImportSummary:
    import_batch_id: int
    framework_rows: int
    mapping_rows: int
    mapping_rows_linked_to_pds: int
    invalid_mapping_codes: int


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "none" else text


def _code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = _clean(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _normalise_header(value: Any) -> str:
    return _clean(value).lower().replace(" ", "_")


def _find_header_row(rows: list[tuple[Any, ...]], required_headers: set[str]) -> int:
    for index, row in enumerate(rows):
        headers = {_normalise_header(value) for value in row if _clean(value)}
        if required_headers <= headers:
            return index
    raise ValueError(f"Could not find required headers: {', '.join(sorted(required_headers))}")


def _row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        headers[index]: row[index] if index < len(row) else None
        for index in range(len(headers))
        if headers[index]
    }


def _worksheet_records(workbook_path: Path, sheet_name: str, required_headers: set[str]) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = _find_header_row(rows, required_headers)
    headers = [_normalise_header(value) for value in rows[header_index]]
    records = []
    for row in rows[header_index + 1:]:
        record = _row_dict(headers, row)
        if any(_clean(value) for value in record.values()):
            records.append(record)
    workbook.close()
    return records


def _find_position_description_ids(connection: sqlite3.Connection, pd_id: str) -> list[int | None]:
    if not pd_id:
        return [None]
    rows = connection.execute(
        """SELECT id FROM position_descriptions
           WHERE position_description_no LIKE ?
              OR source_filename LIKE ?
           ORDER BY position_description_no, source_filename, id""",
        (f"{pd_id}-%", f"{pd_id}-%"),
    ).fetchall()
    return [int(row["id"]) for row in rows] if rows else [None]


def import_job_family_workbook(
    connection: sqlite3.Connection,
    workbook_path: Path,
    *,
    framework_sheet: str = FRAMEWORK_SHEET,
    mapping_sheet: str = MAPPING_SHEET,
) -> JobFamilyImportSummary:
    framework_records = _worksheet_records(
        workbook_path,
        framework_sheet,
        {"code", "name", "level", "parent_code", "main_definition"},
    )
    mapping_records = _worksheet_records(
        workbook_path,
        mapping_sheet,
        {"pd_id", "position_title", "primary_mapping_code", "mapping_validated"},
    )

    framework_rows = [
        {
            "code": _code(record.get("code")),
            "name": _clean(record.get("name")),
            "level": _clean(record.get("level")),
            "parent_code": _code(record.get("parent_code")),
            "main_definition": _clean(record.get("main_definition")),
            "supp_definition_1": _clean(record.get("supp_definition_1")),
            "supp_definition_2": _clean(record.get("supp_definition_2")),
            "exclusions": _clean(record.get("exclusions")),
        }
        for record in framework_records
    ]
    framework_rows = [row for row in framework_rows if row["code"] and row["name"]]
    valid_codes = {row["code"] for row in framework_rows}

    with connection:
        connection.execute("UPDATE job_family_import_batches SET is_active = 0")
        cursor = connection.execute(
            """INSERT INTO job_family_import_batches(
                source_filename, source_path, framework_sheet, mapping_sheet,
                framework_rows, mapping_rows, is_active
            ) VALUES (?, ?, ?, ?, 0, 0, 1)""",
            (workbook_path.name, str(workbook_path), framework_sheet, mapping_sheet),
        )
        batch_id = int(cursor.lastrowid)
        connection.executemany(
            """INSERT INTO job_family_entries(
                import_batch_id, code, name, level, parent_code, main_definition,
                supp_definition_1, supp_definition_2, exclusions, sequence
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    batch_id,
                    row["code"],
                    row["name"],
                    row["level"],
                    row["parent_code"],
                    row["main_definition"],
                    row["supp_definition_1"],
                    row["supp_definition_2"],
                    row["exclusions"],
                    sequence,
                )
                for sequence, row in enumerate(framework_rows, 1)
            ],
        )

        mapping_rows = []
        for record in mapping_records:
            pd_id = _code(record.get("pd_id"))
            matching_position_ids = _find_position_description_ids(connection, pd_id)
            slots = [
                (
                    1,
                    _code(record.get("primary_mapping_code")),
                    _clean(record.get("primary_mapping_name")),
                    _clean(record.get("primary_mapping_level")),
                ),
                (
                    2,
                    _code(record.get("secondary_mapping_code")),
                    _clean(record.get("secondary_mapping_name")),
                    _clean(record.get("secondary_mapping_level")),
                ),
                (
                    3,
                    _code(record.get("secondary_2_mapping_code")),
                    _clean(record.get("secondary_2_mapping_name")),
                    _clean(record.get("secondary_2_mapping_level")),
                ),
            ]
            for position_description_id in matching_position_ids:
                base = {
                    "pd_id": pd_id,
                    "position_description_id": position_description_id,
                    "position_title": _clean(record.get("position_title")),
                    "position_grade": _clean(record.get("position_grade")),
                    "mapping_validated": _clean(record.get("mapping_validated")),
                    "validation_notes": _clean(record.get("validation_notes")),
                }
                for rank, mapping_code, mapping_name, mapping_level in slots:
                    if not mapping_code and not mapping_name:
                        continue
                    mapping_rows.append({
                        **base,
                        "mapping_rank": rank,
                        "mapping_code": mapping_code,
                        "mapping_name": mapping_name,
                        "mapping_level": mapping_level,
                        "framework_code_valid": 1 if mapping_code in valid_codes else 0,
                    })
        connection.executemany(
            """INSERT INTO pd_job_family_mappings(
                import_batch_id, position_description_id, pd_id, position_title,
                position_grade, mapping_rank, mapping_code, mapping_name, mapping_level,
                mapping_validated, validation_notes, framework_code_valid
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    batch_id,
                    row["position_description_id"],
                    row["pd_id"],
                    row["position_title"],
                    row["position_grade"],
                    row["mapping_rank"],
                    row["mapping_code"],
                    row["mapping_name"],
                    row["mapping_level"],
                    row["mapping_validated"],
                    row["validation_notes"],
                    row["framework_code_valid"],
                )
                for row in mapping_rows
            ],
        )
        connection.execute(
            """UPDATE job_family_import_batches
               SET framework_rows = ?, mapping_rows = ?
               WHERE id = ?""",
            (len(framework_rows), len(mapping_rows), batch_id),
        )

    return JobFamilyImportSummary(
        import_batch_id=batch_id,
        framework_rows=len(framework_rows),
        mapping_rows=len(mapping_rows),
        mapping_rows_linked_to_pds=sum(1 for row in mapping_rows if row["position_description_id"]),
        invalid_mapping_codes=sum(1 for row in mapping_rows if not row["framework_code_valid"]),
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Import job family framework and PD mappings from Excel")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--database", type=Path, default=default_database_path())
    parser.add_argument("--framework-sheet", default=FRAMEWORK_SHEET)
    parser.add_argument("--mapping-sheet", default=MAPPING_SHEET)
    args = parser.parse_args()

    connection = connect_database(args.database)
    try:
        initialise_database(connection)
        summary = import_job_family_workbook(
            connection,
            args.workbook,
            framework_sheet=args.framework_sheet,
            mapping_sheet=args.mapping_sheet,
        )
        statistics = database_statistics(connection)
    finally:
        connection.close()
    print(json.dumps({
        "import_batch_id": summary.import_batch_id,
        "framework_rows": summary.framework_rows,
        "mapping_rows": summary.mapping_rows,
        "mapping_rows_linked_to_pds": summary.mapping_rows_linked_to_pds,
        "invalid_mapping_codes": summary.invalid_mapping_codes,
        "database_statistics": statistics,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
from pd_extractor.config import default_database_path
