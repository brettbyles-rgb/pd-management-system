from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from .database import database_backend, database_object_exists
from .job_family_import import _clean, _code, _worksheet_records


FRAMEWORK_SHEET = "job_family_power_query"
MAPPING_SHEET = "job_mapping_input"
ADJACENCY_SHEET = "job_family_adjacency"
INSTRUCTIONS_SHEET = "README"

FRAMEWORK_HEADERS = (
    "code",
    "name",
    "level",
    "parent_code",
    "main_definition",
    "supp_definition_1",
    "supp_definition_2",
    "exclusions",
)
MAPPING_HEADERS = (
    "pd_id",
    "position_title",
    "position_grade",
    "primary_mapping_code",
    "primary_mapping_name",
    "primary_mapping_level",
    "secondary_mapping_code",
    "secondary_mapping_name",
    "secondary_mapping_level",
    "secondary_2_mapping_code",
    "secondary_2_mapping_name",
    "secondary_2_mapping_level",
    "mapping_validated",
    "validation_notes",
)
MAX_WORKBOOK_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 1000


@dataclass(frozen=True)
class ReferenceWorkbookData:
    framework: list[dict[str, str]]
    mappings: list[dict[str, str]]
    adjacency: list[dict[str, Any]]
    family_codes: list[str]


def _rows(connection: Any, sql: str, parameters: object = ()) -> list[dict[str, Any]]:
    return [dict(row) for row in connection.execute(sql, parameters).fetchall()]


def _active_framework(connection: Any) -> list[dict[str, Any]]:
    return _rows(
        connection,
        """SELECT code, name, level, parent_code, main_definition,
                  supp_definition_1, supp_definition_2, exclusions
           FROM active_job_family_entries ORDER BY sequence, code""",
    )


def _effective_mapping_rows(connection: Any) -> list[dict[str, Any]]:
    return _rows(
        connection,
        """SELECT pd_id, position_title, position_grade, mapping_rank,
                  mapping_code, mapping_name, mapping_level,
                  mapping_validated, validation_notes
           FROM active_pd_job_family_mappings
           ORDER BY pd_id, position_description_id, mapping_rank, id""",
    )


def _active_adjacency(connection: Any) -> dict[tuple[str, str], int]:
    if database_object_exists(connection, "active_job_family_adjacency"):
        active = {
            (str(row["source_family_code"]), str(row["target_family_code"])): int(row["tier"])
            for row in connection.execute(
                "SELECT source_family_code, target_family_code, tier "
                "FROM active_job_family_adjacency"
            )
        }
        if active:
            return active
    if database_object_exists(connection, "constellation_family_adjacency"):
        version = connection.execute(
            "SELECT algorithm_version FROM constellation_algorithm_versions "
            "ORDER BY created_at_utc DESC LIMIT 1"
        ).fetchone()
        if version:
            return {
                (str(row["source_family_code"]), str(row["target_family_code"])): int(row["tier"])
                for row in connection.execute(
                    "SELECT source_family_code, target_family_code, tier "
                    "FROM constellation_family_adjacency WHERE algorithm_version = ?",
                    (version["algorithm_version"],),
                )
            }
    return {}


def _style_header(row: Any) -> None:
    fill = PatternFill("solid", fgColor="082646")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)


def _append_safe(worksheet: Any, values: list[Any]) -> None:
    """Write database text as literal Excel values rather than executable formulas."""
    worksheet.append(values)
    for cell, value in zip(worksheet[worksheet.max_row], values):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"


def _autosize(worksheet: Any, *, maximum: int = 48) -> None:
    for column in worksheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=0)
        worksheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), maximum)


def export_reference_workbook(connection: Any) -> bytes:
    framework = _active_framework(connection)
    mapping_rows = _effective_mapping_rows(connection)
    adjacency = _active_adjacency(connection)
    families = [
        row for row in framework if str(row.get("level") or "").strip().lower() == "job family"
    ]

    workbook = Workbook()
    readme = workbook.active
    readme.title = INSTRUCTIONS_SHEET
    readme.append(["PD Management System — governed reference data workbook"])
    readme.append(["Exported UTC", datetime.now(timezone.utc).isoformat()])
    readme.append([])
    readme.append(["Sheet", "Purpose", "Rules"])
    readme.append([FRAMEWORK_SHEET, "Job-family hierarchy", "Codes must be unique; every parent code must exist."])
    readme.append([MAPPING_SHEET, "Role-to-family mappings", "Mapping codes must exist in the framework; maximum three per role."])
    readme.append([ADJACENCY_SHEET, "Directed family adjacency matrix", "Every cell must be 0, 1 or 2; diagonal cells must be 0."])
    readme.append([])
    readme.append(["Workflow", "Download → edit → validate preview → explicitly apply. Never alter sheet names or header rows."])
    readme["A1"].font = Font(size=15, bold=True, color="082646")
    _style_header(readme[4])
    _autosize(readme, maximum=90)

    framework_sheet = workbook.create_sheet(FRAMEWORK_SHEET)
    framework_sheet.append(list(FRAMEWORK_HEADERS))
    for row in framework:
        _append_safe(framework_sheet, [row.get(header, "") for header in FRAMEWORK_HEADERS])
    _style_header(framework_sheet[1])
    framework_sheet.freeze_panes = "A2"
    framework_sheet.auto_filter.ref = framework_sheet.dimensions
    _autosize(framework_sheet)

    mapping_sheet = workbook.create_sheet(MAPPING_SHEET)
    mapping_sheet.append(list(MAPPING_HEADERS))
    grouped: dict[str, dict[int, dict[str, Any]]] = {}
    role_details: dict[str, tuple[str, str]] = {}
    notes: dict[str, tuple[str, str]] = {}
    for row in mapping_rows:
        key = str(row.get("pd_id") or "")
        grouped.setdefault(key, {})[int(row.get("mapping_rank") or 1)] = row
        role_details[key] = (
            str(row.get("position_title") or ""),
            str(row.get("position_grade") or ""),
        )
        notes[key] = (
            str(row.get("mapping_validated") or ""),
            str(row.get("validation_notes") or ""),
        )
    for key in sorted(grouped):
        ranked = grouped[key]
        values: list[Any] = [key, *role_details[key]]
        for rank in (1, 2, 3):
            row = ranked.get(rank, {})
            values.extend(
                [row.get("mapping_code", ""), row.get("mapping_name", ""), row.get("mapping_level", "")]
            )
        values.extend(notes[key])
        _append_safe(mapping_sheet, values)
    _style_header(mapping_sheet[1])
    mapping_sheet.freeze_panes = "A2"
    mapping_sheet.auto_filter.ref = mapping_sheet.dimensions
    _autosize(mapping_sheet)

    adjacency_sheet = workbook.create_sheet(ADJACENCY_SHEET)
    family_codes = [str(row["code"]) for row in families]
    family_names = {str(row["code"]): str(row["name"]) for row in families}
    adjacency_sheet.append(["source_family_code", "source_family_name", *family_codes])
    adjacency_sheet.append(["target family name →", "", *[family_names[code] for code in family_codes]])
    for source in family_codes:
        _append_safe(
            adjacency_sheet,
            [
                source,
                family_names[source],
                *[
                    adjacency.get((source, target), 0 if source == target else "")
                    for target in family_codes
                ],
            ],
        )
    _style_header(adjacency_sheet[1])
    adjacency_sheet.freeze_panes = "C3"
    adjacency_sheet.column_dimensions["A"].width = 22
    adjacency_sheet.column_dimensions["B"].width = 34
    for column_index in range(3, 3 + len(family_codes)):
        adjacency_sheet.column_dimensions[adjacency_sheet.cell(1, column_index).column_letter].width = 15
    if family_codes:
        validation = DataValidation(type="list", formula1='"0,1,2"', allow_blank=False)
        adjacency_sheet.add_data_validation(validation)
        validation.add(
            f"C3:{adjacency_sheet.cell(2 + len(family_codes), 2 + len(family_codes)).coordinate}"
        )
        adjacency_sheet.conditional_formatting.add(
            f"C3:{adjacency_sheet.cell(2 + len(family_codes), 2 + len(family_codes)).coordinate}",
            CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="E2F0D9")),
        )

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _parse_framework(path: Path) -> list[dict[str, str]]:
    records = _worksheet_records(
        path,
        FRAMEWORK_SHEET,
        {"code", "name", "level", "parent_code", "main_definition"},
    )
    return [
        {
            "code": _code(row.get("code")),
            "name": _clean(row.get("name")),
            "level": _clean(row.get("level")),
            "parent_code": _code(row.get("parent_code")),
            "main_definition": _clean(row.get("main_definition")),
            "supp_definition_1": _clean(row.get("supp_definition_1")),
            "supp_definition_2": _clean(row.get("supp_definition_2")),
            "exclusions": _clean(row.get("exclusions")),
        }
        for row in records
        if _code(row.get("code")) or _clean(row.get("name"))
    ]


def _parse_mappings(path: Path) -> list[dict[str, str]]:
    records = _worksheet_records(
        path,
        MAPPING_SHEET,
        {"pd_id", "position_title", "primary_mapping_code", "mapping_validated"},
    )
    return [
        {header: _code(row.get(header)) if header.endswith("_code") or header == "pd_id" else _clean(row.get(header))
         for header in MAPPING_HEADERS}
        for row in records
        if _code(row.get("pd_id")) or _clean(row.get("position_title"))
    ]


def _parse_adjacency(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if ADJACENCY_SHEET not in workbook.sheetnames:
            raise ValueError(f"Workbook does not contain required sheet: {ADJACENCY_SHEET}")
        worksheet = workbook[ADJACENCY_SHEET]
        header = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        target_codes = [_code(value) for value in header[2:] if _code(value)]
        entries: list[dict[str, Any]] = []
        for row in worksheet.iter_rows(min_row=3, values_only=True):
            source = _code(row[0] if row else None)
            if not source:
                continue
            for index, target in enumerate(target_codes, start=2):
                value = row[index] if index < len(row) else None
                tier: Any = value
                if isinstance(value, float) and value.is_integer():
                    tier = int(value)
                elif isinstance(value, str) and value.strip().isdigit():
                    tier = int(value.strip())
                entries.append(
                    {"source_family_code": source, "target_family_code": target, "tier": tier, "rationale": ""}
                )
        return entries, target_codes
    finally:
        workbook.close()


def parse_reference_workbook(path: Path) -> ReferenceWorkbookData:
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook exceeds the 20 MB limit")
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("Workbook archive contains too many files")
        if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("Workbook expands beyond the 100 MB safety limit")
    framework = _parse_framework(path)
    mappings = _parse_mappings(path)
    adjacency, target_codes = _parse_adjacency(path)
    return ReferenceWorkbookData(framework, mappings, adjacency, target_codes)


def validate_reference_workbook(connection: Any, path: Path) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        data = parse_reference_workbook(path)
    except (BadZipFile, InvalidFileException, OSError, StopIteration, ValueError) as error:
        return {"valid": False, "errors": [str(error)], "warnings": [], "counts": {}}

    codes = [row["code"] for row in data.framework]
    code_set = set(codes)
    duplicates = sorted({code for code in codes if codes.count(code) > 1})
    if duplicates:
        errors.append("Duplicate framework codes: " + ", ".join(duplicates[:20]))
    for row in data.framework:
        if not row["code"] or not row["name"] or not row["level"]:
            errors.append("Every framework row requires code, name and level")
            break
        if row["parent_code"] and row["parent_code"] not in code_set:
            errors.append(f"Unknown parent code {row['parent_code']} for {row['code']}")
    parent_by_code = {row["code"]: row["parent_code"] for row in data.framework}
    for code in codes:
        seen: set[str] = set()
        current = code
        while current:
            if current in seen:
                errors.append(f"Framework hierarchy contains a cycle involving {code}")
                break
            seen.add(current)
            current = parent_by_code.get(current, "")

    family_codes = [
        row["code"] for row in data.framework if row["level"].strip().lower() == "job family"
    ]
    if not family_codes:
        errors.append("The framework contains no rows with level 'Job Family'")
    if data.family_codes != family_codes:
        errors.append("Adjacency column codes must exactly match the Job Family rows in framework order")
    sources = [row["source_family_code"] for row in data.adjacency[::max(len(data.family_codes), 1)]]
    if sources != family_codes:
        errors.append("Adjacency row codes must exactly match the Job Family rows in framework order")
    expected_cells = len(family_codes) ** 2
    if len(data.adjacency) != expected_cells:
        errors.append(f"Adjacency matrix must contain {expected_cells} cells; found {len(data.adjacency)}")
    matrix: dict[tuple[str, str], Any] = {}
    for row in data.adjacency:
        key = (row["source_family_code"], row["target_family_code"])
        if key in matrix:
            errors.append(f"Duplicate adjacency cell {key[0]} → {key[1]}")
        matrix[key] = row["tier"]
        if row["tier"] not in (0, 1, 2):
            errors.append(f"Adjacency cell {key[0]} → {key[1]} must be 0, 1 or 2")
        if key[0] == key[1] and row["tier"] != 0:
            errors.append(f"Adjacency diagonal {key[0]} → {key[1]} must be 0")
    asymmetric = [
        f"{source} ↔ {target}"
        for (source, target), tier in matrix.items()
        if source < target and matrix.get((target, source)) in (0, 1, 2) and matrix[(target, source)] != tier
    ]
    if asymmetric:
        warnings.append("Directed adjacency values differ for: " + ", ".join(asymmetric[:20]))

    mapped_codes: list[str] = []
    mapping_ids = [row["pd_id"] for row in data.mappings if row["pd_id"]]
    duplicate_mapping_ids = sorted(
        {pd_id for pd_id in mapping_ids if mapping_ids.count(pd_id) > 1}
    )
    if duplicate_mapping_ids:
        errors.append("Duplicate role mapping rows: " + ", ".join(duplicate_mapping_ids[:20]))
    for row in data.mappings:
        if not row["pd_id"]:
            errors.append("Every mapping row requires pd_id")
        if not row["primary_mapping_code"]:
            errors.append(f"Role {row['pd_id'] or '(blank)'} requires a primary mapping code")
        role_codes = [
            row[field]
            for field in ("primary_mapping_code", "secondary_mapping_code", "secondary_2_mapping_code")
            if row[field]
        ]
        if len(role_codes) != len(set(role_codes)):
            errors.append(f"Role {row['pd_id']} contains the same mapping code more than once")
        for field in ("primary_mapping_code", "secondary_mapping_code", "secondary_2_mapping_code"):
            value = row[field]
            if value:
                mapped_codes.append(value)
                if value not in code_set:
                    errors.append(f"Role {row['pd_id']} uses unknown mapping code {value}")

    database_pd_ids = {
        str(row["position_description_no"] or "").split("-", 1)[0]
        for row in connection.execute("SELECT position_description_no FROM position_descriptions")
        if row["position_description_no"]
    }
    unlinked = sorted({row["pd_id"] for row in data.mappings if row["pd_id"] not in database_pd_ids})
    if unlinked:
        warnings.append(f"{len(unlinked)} workbook role IDs do not currently match a position description")

    current_framework = _active_framework(connection)
    current_mappings = _effective_mapping_rows(connection)
    current_adjacency = _active_adjacency(connection)
    incoming_framework = {row["code"]: tuple(row.get(header, "") for header in FRAMEWORK_HEADERS[1:]) for row in data.framework}
    existing_framework = {
        str(row["code"]): tuple(str(row.get(header) or "") for header in FRAMEWORK_HEADERS[1:])
        for row in current_framework
    }
    incoming_mapping_set = {
        (row["pd_id"], rank, row[field])
        for row in data.mappings
        for rank, field in enumerate(
            ("primary_mapping_code", "secondary_mapping_code", "secondary_2_mapping_code"),
            start=1,
        )
        if row[field]
    }
    existing_mapping_set = {
        (str(row.get("pd_id") or ""), int(row.get("mapping_rank") or 1), str(row.get("mapping_code") or ""))
        for row in current_mappings
        if row.get("mapping_code")
    }
    incoming_adjacency = {
        (row["source_family_code"], row["target_family_code"]): row["tier"]
        for row in data.adjacency
    }
    removed_framework = existing_framework.keys() - incoming_framework.keys()
    removed_mappings = existing_mapping_set - incoming_mapping_set
    if removed_framework:
        warnings.append(f"Applying this workbook removes {len(removed_framework)} framework codes")
    if removed_mappings:
        warnings.append(f"Applying this workbook removes {len(removed_mappings)} current mapping assignments")
    return {
        "valid": not errors,
        "errors": list(dict.fromkeys(errors)),
        "warnings": warnings,
        "counts": {
            "framework_rows": len(data.framework),
            "mapping_role_rows": len(data.mappings),
            "mapping_assignments": len(mapped_codes),
            "job_families": len(family_codes),
            "adjacency_cells": len(data.adjacency),
        },
        "current": {
            "framework_rows": len(current_framework),
            "mapping_assignments": len(existing_mapping_set),
            "adjacency_cells": len(current_adjacency),
        },
        "changes": {
            "framework_added": len(incoming_framework.keys() - existing_framework.keys()),
            "framework_removed": len(existing_framework.keys() - incoming_framework.keys()),
            "framework_changed": sum(
                1
                for code in incoming_framework.keys() & existing_framework.keys()
                if incoming_framework[code] != existing_framework[code]
            ),
            "mapping_assignments_added": len(incoming_mapping_set - existing_mapping_set),
            "mapping_assignments_removed": len(existing_mapping_set - incoming_mapping_set),
            "adjacency_cells_added": len(incoming_adjacency.keys() - current_adjacency.keys()),
            "adjacency_cells_removed": len(current_adjacency.keys() - incoming_adjacency.keys()),
            "adjacency_cells_changed": sum(
                1
                for key in incoming_adjacency.keys() & current_adjacency.keys()
                if incoming_adjacency[key] != current_adjacency[key]
            ),
        },
    }


def _next_id(connection: Any, table: str) -> int:
    return int(connection.execute(f"SELECT COALESCE(MAX(id), 0) + 1 AS next_id FROM {table}").fetchone()["next_id"])


def _pd_ids_by_base(connection: Any) -> dict[str, list[int]]:
    output: dict[str, list[int]] = {}
    for row in connection.execute(
        "SELECT id, position_description_no, source_filename FROM position_descriptions ORDER BY id"
    ):
        label = str(row["position_description_no"] or row["source_filename"] or "")
        base = label.split("-", 1)[0]
        output.setdefault(base, []).append(int(row["id"]))
    return output


def apply_reference_workbook(connection: Any, path: Path) -> dict[str, Any]:
    validation = validate_reference_workbook(connection, path)
    if not validation["valid"]:
        raise ValueError("; ".join(validation["errors"]))
    data = parse_reference_workbook(path)
    imported_at = datetime.now(timezone.utc).isoformat()
    by_code = {row["code"]: row for row in data.framework}
    pd_ids = _pd_ids_by_base(connection)

    with connection:
        if database_backend(connection) == "postgresql":
            connection.execute(
                "LOCK TABLE job_family_import_batches, job_family_entries, "
                "pd_job_family_mappings, job_family_adjacency_entries IN EXCLUSIVE MODE"
            )
        batch_id = _next_id(connection, "job_family_import_batches")
        entry_id = _next_id(connection, "job_family_entries")
        mapping_id = _next_id(connection, "pd_job_family_mappings")
        mapping_rows: list[tuple[Any, ...]] = []
        for source in data.mappings:
            targets = pd_ids.get(source["pd_id"], [None])
            ranked = (
                (1, "primary_mapping_code"),
                (2, "secondary_mapping_code"),
                (3, "secondary_2_mapping_code"),
            )
            for position_id in targets:
                for rank, code_key in ranked:
                    code = source[code_key]
                    if not code:
                        continue
                    framework = by_code[code]
                    mapping_rows.append(
                        (
                            mapping_id,
                            batch_id,
                            position_id,
                            source["pd_id"],
                            source["position_title"],
                            source["position_grade"],
                            rank,
                            code,
                            framework["name"],
                            framework["level"],
                            source["mapping_validated"],
                            source["validation_notes"],
                            1,
                        )
                    )
                    mapping_id += 1
        connection.execute("UPDATE job_family_import_batches SET is_active = 0")
        connection.execute(
            """INSERT INTO job_family_import_batches(
                   id, source_filename, source_path, framework_sheet, mapping_sheet,
                   framework_rows, mapping_rows, is_active, imported_at
               ) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)""",
            (
                batch_id,
                path.name,
                None,
                FRAMEWORK_SHEET,
                MAPPING_SHEET,
                len(data.framework),
                len(mapping_rows),
                imported_at,
            ),
        )
        entries = []
        for sequence, row in enumerate(data.framework, 1):
            entries.append(
                (
                    entry_id,
                    batch_id,
                    row["code"],
                    row["name"],
                    row["level"],
                    row["parent_code"] or None,
                    row["main_definition"],
                    row["supp_definition_1"],
                    row["supp_definition_2"],
                    row["exclusions"],
                    sequence,
                )
            )
            entry_id += 1
        connection.executemany(
            """INSERT INTO job_family_entries(
                   id, import_batch_id, code, name, level, parent_code,
                   main_definition, supp_definition_1, supp_definition_2,
                   exclusions, sequence
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            entries,
        )
        connection.executemany(
            """INSERT INTO pd_job_family_mappings(
                   id, import_batch_id, position_description_id, pd_id,
                   position_title, position_grade, mapping_rank, mapping_code,
                   mapping_name, mapping_level, mapping_validated,
                   validation_notes, framework_code_valid
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            mapping_rows,
        )
        connection.executemany(
            """INSERT INTO job_family_adjacency_entries(
                   import_batch_id, source_family_code, target_family_code, tier, rationale
               ) VALUES (?, ?, ?, ?, ?)""",
            [
                (
                    batch_id,
                    row["source_family_code"],
                    row["target_family_code"],
                    row["tier"],
                    row["rationale"],
                )
                for row in data.adjacency
            ],
        )
        connection.execute("DELETE FROM pd_assigned_job_family_mappings")

    return {
        "import_batch_id": batch_id,
        "imported_at": imported_at,
        **validation["counts"],
        "linked_position_descriptions": sum(1 for row in mapping_rows if row[2] is not None),
    }
