from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from pd_extractor.activities import prompts


ROOT = Path(r"C:\Users\BrettB\Documents\pd_management_system")
OUTPUT = ROOT / "output" / "activity-summarising-method-for-claude.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, bold=True)
        set_cell_shading(header_cells[index], "F2F4F7")
        if widths:
            header_cells[index].width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph()


def add_code_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    for line in text.strip("\n").splitlines():
        run = paragraph.add_run(line + "\n")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8)


def main() -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Activity Summarising and Clustering Method")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 77, 120)

    subtitle = doc.add_paragraph()
    subtitle.add_run("Current algorithm and prompts for Claude review").italic = True

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This brief documents the current activity summarising pipeline used to convert position description "
        "key accountabilities into reusable, clustered work activities. The downstream objective is to use "
        "shared clustered activities as linkage points between roles for a career pathways explorer."
    )

    doc.add_heading("What We Want Claude To Review", level=1)
    add_bullets(
        doc,
        [
            "Whether the extracted activity labels are at the right level of abstraction for comparing roles.",
            "Whether the clustering/vocabulary approach creates useful common activities without becoming too generic.",
            "Whether the assignment stage should choose fewer, more discriminating activities for career-path linkage.",
            "Whether boilerplate exclusion is too narrow, too broad, or should be replaced by a weighted approach.",
            "Whether the prompts should distinguish core, secondary, enabling, and generic activities.",
        ],
    )

    doc.add_heading("Current Pipeline Overview", level=1)
    add_table(
        doc,
        ["Stage", "Input", "Method", "Output"],
        [
            [
                "0. Load and filter",
                "PD key accountabilities from SQLite",
                "Clean list items and exclude known boilerplate accountabilities using the same rules as the mapping/embedding process.",
                "Role records containing role-specific accountabilities plus an audit list of excluded items.",
            ],
            [
                "1. Extract candidate activities",
                "Filtered accountabilities per role",
                "Kimi extracts generic activity labels from each role. It may split bundled accountabilities into multiple activities.",
                "1_extracted.jsonl: one record per PD, containing candidate activities with source accountability indexes.",
            ],
            [
                "2. Build shared vocabulary",
                "All extracted match_label phrases",
                "Count repeated phrases, keep phrases seen at least twice, embed them locally, cluster near-duplicates, ask Kimi to name canonical activities.",
                "2_vocabulary.json and 2_cluster_members.json: canonical activities and their raw member phrases.",
            ],
            [
                "3. Assign clustered activities to roles",
                "Role text plus fixed vocabulary shortlist",
                "Use local embedding similarity to shortlist likely activities, then ask Kimi to choose 6-14 activity IDs for each role.",
                "3_assigned.jsonl: one record per PD with ranked canonical activity IDs.",
            ],
            [
                "4. Export for analysis",
                "Assigned roles plus vocabulary",
                "Flatten to role x activity links and a role-activity matrix.",
                "Excel workbook for commonality and career-path analysis.",
            ],
        ],
        widths=[1.1, 1.45, 2.05, 1.9],
    )

    doc.add_heading("Current 200-Role Test Metrics", level=1)
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Roles processed", "200"],
            ["Extracted candidate activity rows", "2,843"],
            ["Excluded boilerplate accountabilities", "496"],
            ["Canonical vocabulary rows after clustering", "102"],
            ["Assigned role x clustered activity rows", "2,399"],
            ["Assigned clustered activities per role", "Minimum 7; median 12; maximum 14"],
            ["Assignment retries", "7"],
            ["Assignment fallbacks", "0"],
        ],
        widths=[3.1, 3.0],
    )

    doc.add_heading("Boilerplate Exclusion", level=1)
    doc.add_paragraph(
        "Before activity extraction, the loader removes accountabilities identified as boilerplate. Removed items are not sent "
        "to Kimi, but are retained in the output for audit. This is intentionally aligned with the existing PD mapping "
        "embedding process."
    )
    add_table(
        doc,
        ["Exclusion type", "Current examples or patterns"],
        [
            ["Mandatory values/safety", "Reflect TAFE NSW's values; policies and procedures; safe, healthy and inclusive work environment."],
            ["Customer-centred decision boilerplate", "Place the customer at the centre of all decision making."],
            ["Safety leadership", "Safety excellence, safety leadership, safe workplace, safety systems and procedures."],
            ["High-performance team", "High-performance team; core values of integrity, collaboration, excellence; leadership, support and feedback."],
            ["Performance planning", "Individual performance management and development plans; regular review of meaningful individual performance."],
        ],
        widths=[2.0, 4.1],
    )

    doc.add_heading("Stage 1: Candidate Activity Extraction", level=1)
    doc.add_paragraph(
        "The extraction stage asks Kimi to produce candidate activity phrases for each role. These are not yet the final "
        "shared vocabulary. Similar or duplicate labels are expected at this stage."
    )
    doc.add_heading("Extraction System Prompt", level=2)
    add_code_block(doc, prompts.EXTRACT_SYSTEM)
    doc.add_heading("Extraction User Prompt Template", level=2)
    add_code_block(doc, prompts.EXTRACT_USER)
    doc.add_paragraph("Implementation notes:")
    add_bullets(
        doc,
        [
            "Kimi model: kimi-k3 via OpenAI-compatible API.",
            "Reasoning effort currently set to high in the launchers.",
            "Extraction max completion tokens: 4,000.",
            "Blank or invalid extraction responses are retried up to two times.",
            "Stage output records source accountability indexes in the from field.",
        ],
    )

    doc.add_heading("Stage 2: Vocabulary and Clustering", level=1)
    doc.add_paragraph(
        "The clustering stage creates the shared activity vocabulary. It counts extracted match_label phrases, retains phrases "
        "seen at least twice, embeds the retained phrases with a local TF-IDF/SVD embedder, clusters them using agglomerative "
        "clustering, and asks Kimi to name each cluster. A cluster may be split into two or rarely three canonical activities "
        "if Kimi judges the embedding cluster to contain multiple meanings."
    )
    doc.add_heading("Clustering System Prompt", level=2)
    add_code_block(doc, prompts.CLUSTER_SYSTEM)
    doc.add_heading("Clustering User Prompt Template", level=2)
    add_code_block(doc, prompts.CLUSTER_USER)
    doc.add_paragraph("Implementation notes:")
    add_bullets(
        doc,
        [
            "Only phrases with raw count >= 2 are included in the clustering stage.",
            "Default target cluster count is 280, but the actual count is capped by sample size and phrase count.",
            "For the 200-role test, the output vocabulary contains 102 canonical activities.",
            "The local embedder is TF-IDF with unigrams/bigrams, TruncatedSVD, and vector normalization.",
            "Cluster naming max completion tokens: 800.",
        ],
    )

    doc.add_heading("Stage 3: Assign Clustered Activities Back To Roles", level=1)
    doc.add_paragraph(
        "The assignment stage is the main input for career-path commonality analysis. For each role, the code embeds the role "
        "text and vocabulary labels, shortlists the top 60 vocabulary activities by similarity, then asks Kimi to choose the "
        "6-14 activity IDs that genuinely describe the role's day-to-day work."
    )
    doc.add_heading("Assignment System Prompt", level=2)
    add_code_block(doc, prompts.ASSIGN_SYSTEM)
    doc.add_heading("Assignment User Prompt Template", level=2)
    add_code_block(doc, prompts.ASSIGN_USER)
    doc.add_paragraph("Implementation notes:")
    add_bullets(
        doc,
        [
            "Shortlist size: 60 vocabulary activities per role.",
            "Kimi may only choose IDs from the supplied shortlist.",
            "The role gets between 6 and 14 assigned clustered activities.",
            "Assignment max completion tokens: 1,200.",
            "Assignment runs role-by-role for resilience; retries blank or invalid responses; falls back to the embedding shortlist if needed.",
            "For the current 200-role test, no fallbacks were used.",
        ],
    )

    doc.add_heading("Known Design Choices and Possible Weak Points", level=1)
    add_bullets(
        doc,
        [
            "The extraction prompt tries to split bundled accountabilities, which can create many fine-grained activity labels.",
            "The clustering stage ignores one-off extracted phrases by default, so rare activities may not appear in the vocabulary.",
            "Some non-discriminating activities remain in the vocabulary and assignments, such as performance planning, record keeping, and safe workplace activities.",
            "The assignment prompt asks for 6-14 activities but does not explicitly prioritise career-path signal over general role coverage.",
            "The system currently has a discriminating flag on vocabulary items, but the assignment stage does not require a minimum number of discriminating activities.",
            "The role activity matrix is useful for overlap analysis, but common generic activities may create weak or misleading links unless filtered or weighted.",
        ],
    )

    doc.add_heading("Specific Questions For Claude", level=1)
    add_bullets(
        doc,
        [
            "Should activity extraction target fewer, broader career-relevant activity groups, or keep the current more granular activity style?",
            "Should generic/enabling activities be excluded from the role matrix, down-weighted, or retained as a separate feature type?",
            "Should assignment output distinguish core activities from secondary or enabling activities?",
            "Should the vocabulary include rare but career-significant activities that appear only once in the sample?",
            "Should clustering be adjusted to reduce duplicate canonical activities or improve interpretability?",
            "What prompt changes would make the role activity matrix more useful for career-path linkages?",
        ],
    )

    doc.add_heading("Current Export Used For Analysis", level=1)
    doc.add_paragraph(
        "The most recent Claude analysis workbook is: "
        r"C:\Users\BrettB\Documents\pd_management_system\output\activity-assigned-clustered-role-links-200.xlsx"
    )
    doc.add_paragraph(
        "The key worksheet is Role Activity Matrix: each row is a role; each canonical activity column is a 1/0 flag."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
