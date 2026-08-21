from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def unique_preserve_order(items: list[str]) -> list[str]:
    seen = set()
    out = []
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        out.append(item)
    return out


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    parser = argparse.ArgumentParser(description="Build all-role activity workbook with openpyxl")
    parser.add_argument("--work", type=Path, default=Path("work/activities-revised-200-high"))
    parser.add_argument("--output", type=Path, default=Path("output/activity-assigned-clustered-role-links-revised-high-all.xlsx"))
    args = parser.parse_args()

    work = args.work
    output = args.output
    output.parent.mkdir(parents=True, exist_ok=True)

    extracted_records = read_jsonl(work / "1_extracted.jsonl")
    assigned_records = read_jsonl(work / "3_assigned.jsonl")
    vocabulary = json.loads((work / "2_vocabulary.json").read_text(encoding="utf-8"))
    cluster_members = json.loads((work / "2_cluster_members.json").read_text(encoding="utf-8"))

    vocab_by_id = {item["id"]: item for item in vocabulary}
    extracted_by_pd = {str(item["pd_id"]): item for item in extracted_records}

    clean_assignments = []
    usage = Counter()
    link_count = 0
    for record in assigned_records:
        ids = unique_preserve_order(record.get("activity_ids") or [])
        usage.update(ids)
        link_count += len(ids)
        clean_assignments.append((record, ids))

    wb = Workbook(write_only=False)
    wb.remove(wb.active)

    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    sizes = [len(ids) for _, ids in clean_assignments]
    summary_rows = [
        ["Assigned role count", len(clean_assignments)],
        ["Role x clustered activity rows", link_count],
        ["Canonical vocabulary rows", len(vocabulary)],
        ["Min assigned activities per role", min(sizes) if sizes else 0],
        ["Max assigned activities per role", max(sizes) if sizes else 0],
        ["Retries used", sum(int(record.get("retry_count") or 0) for record, _ in clean_assignments)],
        ["Fallbacks used", sum(1 for record, _ in clean_assignments if record.get("fallback_used"))],
    ]
    for row in summary_rows:
        summary.append(row)

    role_links = wb.create_sheet("Role Activity Links")
    role_links.append(
        [
            "PD ID",
            "PD Number",
            "Title",
            "Activity Rank",
            "Activity ID",
            "Canonical Match Label",
            "Canonical Plain Label",
            "Discriminating",
            "Cluster ID",
            "Raw Count",
            "Cluster Member Phrases",
            "Assignment Gaps",
            "Retry Count",
            "Fallback Used",
        ]
    )
    for record, ids in clean_assignments:
        for rank, activity_id in enumerate(ids, start=1):
            vocab = vocab_by_id.get(activity_id, {})
            role_links.append(
                [
                    record.get("position_description_id") or record.get("pd_id") or "",
                    record.get("position_description_no") or "",
                    record.get("title") or "",
                    rank,
                    activity_id,
                    vocab.get("match_label", ""),
                    vocab.get("plain_label", ""),
                    "Yes" if vocab.get("discriminating", True) else "No",
                    vocab.get("cluster", ""),
                    vocab.get("raw_count", ""),
                    " | ".join(cluster_members.get(activity_id, [])),
                    " | ".join(record.get("gaps") or []),
                    record.get("retry_count") or 0,
                    "Yes" if record.get("fallback_used") else "No",
                ]
            )

    role_summary = wb.create_sheet("Role Summary")
    role_summary.append(
        [
            "PD ID",
            "PD Number",
            "Title",
            "Source Accountability Count",
            "Excluded Boilerplate Count",
            "Extracted Activity Count",
            "Assigned Clustered Activity Count",
            "Retry Count",
            "Fallback Used",
            "Activity IDs",
            "Canonical Match Labels",
            "Canonical Plain Labels",
            "Assignment Gaps",
        ]
    )
    for record, ids in clean_assignments:
        extracted = extracted_by_pd.get(str(record.get("pd_id")), {})
        labels = [vocab_by_id.get(activity_id, {}).get("match_label", "") for activity_id in ids]
        plain_labels = [vocab_by_id.get(activity_id, {}).get("plain_label", "") for activity_id in ids]
        role_summary.append(
            [
                record.get("position_description_id") or record.get("pd_id") or "",
                record.get("position_description_no") or "",
                record.get("title") or "",
                extracted.get("source_accountability_count", ""),
                len(extracted.get("excluded_accountabilities") or []),
                len(extracted.get("activities") or []),
                len(ids),
                record.get("retry_count") or 0,
                "Yes" if record.get("fallback_used") else "No",
                " | ".join(ids),
                " | ".join(label for label in labels if label),
                " | ".join(label for label in plain_labels if label),
                " | ".join(record.get("gaps") or []),
            ]
        )

    vocab_usage = wb.create_sheet("Vocabulary Usage")
    vocab_usage.append(
        [
            "Activity ID",
            "Canonical Match Label",
            "Canonical Plain Label",
            "Discriminating",
            "Cluster ID",
            "Raw Count",
            "Raw Phrase Count",
            "Singleton Phrase Count",
            "Rare Phrase Count",
            "Has Singleton Member",
            "Min Member Count",
            "Max Member Count",
            "Assigned Role Count",
            "Assigned Role Share",
            "Cluster Member Phrases",
        ]
    )
    total_roles = max(1, len(clean_assignments))
    for item in vocabulary:
        activity_id = item["id"]
        vocab_usage.append(
            [
                activity_id,
                item.get("match_label", ""),
                item.get("plain_label", ""),
                "Yes" if item.get("discriminating", True) else "No",
                item.get("cluster", ""),
                item.get("raw_count", ""),
                item.get("raw_phrase_count", ""),
                item.get("singleton_phrase_count", ""),
                item.get("rare_phrase_count", ""),
                "Yes" if item.get("has_singleton_member") else "No",
                item.get("min_member_count", ""),
                item.get("max_member_count", ""),
                usage[activity_id],
                usage[activity_id] / total_roles,
                " | ".join(cluster_members.get(activity_id, [])),
            ]
        )

    matrix = wb.create_sheet("Role Activity Matrix")
    matrix.append(["PD ID", "PD Number", "Title", *[f"{item['id']}: {item.get('match_label', '')}" for item in vocabulary]])
    vocab_ids = [item["id"] for item in vocabulary]
    for record, ids in clean_assignments:
        id_set = set(ids)
        matrix.append(
            [
                record.get("position_description_id") or record.get("pd_id") or "",
                record.get("position_description_no") or "",
                record.get("title") or "",
                *[1 if activity_id in id_set else 0 for activity_id in vocab_ids],
            ]
        )

    for ws in wb.worksheets:
        style_header(ws)
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 18
    for ws in [role_links, role_summary, vocab_usage]:
        for index in range(1, min(ws.max_column, 15) + 1):
            ws.column_dimensions[get_column_letter(index)].width = 18
        ws.column_dimensions["C"].width = 42
    role_links.column_dimensions["K"].width = 58
    role_links.column_dimensions["L"].width = 58
    role_summary.column_dimensions["J"].width = 58
    role_summary.column_dimensions["K"].width = 58
    role_summary.column_dimensions["L"].width = 58
    role_summary.column_dimensions["M"].width = 58
    vocab_usage.column_dimensions["O"].width = 72
    matrix.freeze_panes = "D2"
    matrix.column_dimensions["A"].width = 10
    matrix.column_dimensions["B"].width = 14
    matrix.column_dimensions["C"].width = 42

    wb.save(output)
    print(json.dumps({"output": str(output), "roles": len(clean_assignments), "links": link_count, "vocabulary": len(vocabulary)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
