from pathlib import Path

from openpyxl import Workbook

from pd_extractor import extract_document
from pd_extractor.bulk import run_bulk_import
from pd_extractor.database import (
    assign_job_family_mappings,
    connect_database,
    database_statistics,
    generate_database_report,
    import_document,
    initialise_database,
    job_family_mappings_for_pd,
)
from pd_extractor.embeddings import generate_embeddings
from pd_extractor.intelligence_app import search_pds
from pd_extractor.job_family_import import import_job_family_workbook
from pd_extractor.mapping_suggestions import (
    enrich_mapping_suggestion_hierarchies,
    suggest_mappings_from_similar_roles,
)
from pd_extractor.mapping_text import build_mapping_text_record, generate_mapping_texts
from pd_extractor.similarity import SimilarRole, cosine_similarity, find_similar_roles, search_roles

SAMPLES = Path(__file__).parent / "fixtures"


class FakeEmbedder:
    def encode(self, texts, **kwargs):
        vectors = []
        for text in texts:
            seed = float(len(text) or 1)
            vectors.append([seed, seed / 2, 1.0])
        return vectors


class QueryFakeEmbedder:
    def encode(self, texts, **kwargs):
        return [[1, 0, 0] for _ in texts]


def test_database_import_is_normalized_and_idempotent(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        adr = extract_document(SAMPLES / "10108-01 Alternate Dispute Resolution Specialist - TW9.docx")
        cfo = extract_document(SAMPLES / "10773-02 Chief Financial Officer - PSSE Band 2.docx")
        import_document(connection, adr)
        import_document(connection, cfo)
        import_document(connection, cfo)

        assert connection.execute("SELECT COUNT(*) FROM position_descriptions").fetchone()[0] == 2
        assert connection.execute("SELECT COUNT(*) FROM pd_capabilities").fetchone()[0] == 44
        frameworks = dict(connection.execute("""
            SELECT cf.framework_name, COUNT(pc.id)
            FROM capability_frameworks cf
            JOIN capability_definitions cd ON cd.framework_id = cf.id
            JOIN pd_capabilities pc ON pc.capability_definition_id = cd.id
            GROUP BY cf.id
        """).fetchall())
        assert frameworks["Human resources"] == 2
        assert frameworks["Finance"] == 6
        assert frameworks["NSW Public Sector Capability Framework"] == 36
        assert connection.execute("SELECT COUNT(*) FROM capability_indicators").fetchone()[0] == 14
        assert database_statistics(connection)["integrity"] == "ok"
        readable = connection.execute("""
            SELECT role_title, capability_type, framework, capability_name,
                   capability_code, required_level
            FROM pd_capabilities_readable
            WHERE framework = 'Finance'
            ORDER BY sequence
        """).fetchall()
        assert len(readable) == 6
        assert readable[0]["role_title"] == "Chief Financial Officer"
        assert readable[0]["capability_type"] == "Complementary"
        assert readable[0]["capability_name"] == "Financial Strategy, Governance and Risk Management"
        assert readable[0]["required_level"] == "Level 5"
    finally:
        connection.close()


def test_database_report_contains_imported_content(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        import_document(
            connection,
            extract_document(SAMPLES / "10108-01 Alternate Dispute Resolution Specialist - TW9.docx"),
        )
        report = generate_database_report(connection, tmp_path / "database-summary.html")
        content = report.read_text(encoding="utf-8")
        assert "Alternate Dispute Resolution Specialist" in content
        assert "Human resources" in content
        assert "Position Description Database" in content
    finally:
        connection.close()


def test_bulk_import_writes_outputs_and_database(tmp_path):
    output_dir = tmp_path / "bulk-json"
    database = tmp_path / "bulk.sqlite3"
    report = tmp_path / "bulk-summary.html"

    summary = run_bulk_import(
        SAMPLES,
        output_dir,
        database,
        limit=2,
        overwrite_database=True,
        report=report,
    )

    assert summary.discovered >= 2
    assert summary.attempted == 2
    assert summary.extracted == 2
    assert summary.failed == 0
    assert summary.imported == 2
    assert summary.summary_json.exists()
    assert summary.summary_csv.exists()
    assert len(list(output_dir.glob("*.json"))) == 3
    assert report.exists()

    connection = connect_database(database)
    try:
        assert database_statistics(connection)["position_descriptions"] == 2
    finally:
        connection.close()


def test_job_family_workbook_import_is_versioned_and_validates_codes(tmp_path):
    workbook_path = tmp_path / "mapping.xlsx"
    workbook = Workbook()
    mapping = workbook.active
    mapping.title = "job_mapping_input"
    mapping.append([
        "pd_id", "position_title", "position_grade",
        "primary_mapping_code", "primary_mapping_name", "primary_mapping_level",
        "secondary_mapping_code", "secondary_mapping_name", "secondary_mapping_level",
        "secondary_2_mapping_code", "secondary_2_mapping_name", "secondary_2_mapping_level",
        "mapping_validated", "validation_notes",
    ])
    mapping.append([
        10021, "Facilities Officer", "TW4",
        2000000000, "FACILITIES", "Job Family",
        9999999999, "NOT REAL", "Sub-family",
        None, None, None,
        "Yes", "checked",
    ])
    framework = workbook.create_sheet("job_family_power_query")
    framework.append([])
    framework.append([
        "Code", "Name", "Level", "Parent Code", "main_definition",
        "supp_definition_1", "supp_definition_2", "exclusions",
    ])
    framework.append([
        2000000000, "FACILITIES", "Job Family", None,
        "Facilities roles manage places and physical services.",
        None, None, None,
    ])
    workbook.save(workbook_path)

    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        import_document(
            connection,
            extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx"),
        )
        second_version = extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx")
        second_version["pd_record"]["source_filename"] = "10021-02 Facilities Officer - TW4.docx"
        second_version["role_description_fields"]["position_description_no"] = "10021-02"
        for raw_field in second_version["role_description_fields"]["raw_fields"]:
            if raw_field["field_name"] == "Position Description no":
                raw_field["field_value"] = "10021-02"
        import_document(connection, second_version)
        summary = import_job_family_workbook(connection, workbook_path)
        assert summary.framework_rows == 1
        assert summary.mapping_rows == 4
        assert summary.mapping_rows_linked_to_pds == 4
        assert summary.invalid_mapping_codes == 2
        assert connection.execute("SELECT COUNT(*) FROM active_job_family_entries").fetchone()[0] == 1
        assert connection.execute(
            "SELECT COUNT(*) FROM active_pd_job_family_mappings WHERE framework_code_valid = 1"
        ).fetchone()[0] == 2
        assert connection.execute(
            "SELECT COUNT(DISTINCT position_description_id) FROM active_pd_job_family_mappings"
        ).fetchone()[0] == 2

        second = import_job_family_workbook(connection, workbook_path)
        assert second.import_batch_id != summary.import_batch_id
        assert connection.execute(
            "SELECT COUNT(*) FROM job_family_import_batches WHERE is_active = 1"
        ).fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM active_job_family_entries").fetchone()[0] == 1
    finally:
        connection.close()


def test_mapping_assistant_assigns_up_to_three_codes_and_overrides_workbook_mapping(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        pd_id = import_document(
            connection,
            extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx"),
        )
        connection.execute(
            """INSERT INTO job_family_import_batches(
                   source_filename, framework_sheet, mapping_sheet, is_active
               ) VALUES ('test.xlsx', 'job_family_power_query', 'job_mapping_input', 1)"""
        )
        batch_id = connection.execute(
            "SELECT id FROM job_family_import_batches"
        ).fetchone()["id"]
        entries = [
            (batch_id, "100", "BUSINESS SUPPORT", "Job Family", None, 1),
            (batch_id, "110", "ADMINISTRATION", "Sub-family", "100", 2),
            (batch_id, "111", "OFFICE ADMINISTRATION", "Specialisation", "110", 3),
        ]
        connection.executemany(
            """INSERT INTO job_family_entries(
                   import_batch_id, code, name, level, parent_code, sequence
               ) VALUES (?, ?, ?, ?, ?, ?)""",
            entries,
        )
        connection.execute(
            """INSERT INTO pd_job_family_mappings(
                   import_batch_id, position_description_id, pd_id, mapping_rank,
                   mapping_code, mapping_name, mapping_level, mapping_validated,
                   framework_code_valid
               ) VALUES (?, ?, '10021', 1, '100', 'BUSINESS SUPPORT',
                         'Job Family', 'Yes', 1)""",
            (batch_id, pd_id),
        )
        connection.commit()

        assigned = assign_job_family_mappings(connection, pd_id, ["110", "111"])

        assert [row["mapping_code"] for row in assigned] == ["110", "111"]
        assert job_family_mappings_for_pd(connection, pd_id) == assigned
        assert connection.execute(
            """SELECT COUNT(*) FROM active_pd_job_family_mappings
               WHERE position_description_id = ? AND mapping_code = '100'""",
            (pd_id,),
        ).fetchone()[0] == 0
    finally:
        connection.close()


def test_mapping_text_excludes_boilerplate_and_keeps_audit_trail(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        data = extract_document(SAMPLES / "11358-01 Director Agency Alignment and Integration - PSSE Band 1.docx")
        pd_id = import_document(connection, data)

        record = build_mapping_text_record(connection, pd_id)

        assert "Director Agency Alignment and Integration" in record.full_text
        assert "Monitor delivery of key strategic initiatives" in record.full_text
        assert "Place the customer at the centre of all decision making" not in record.full_text
        assert "Working with Children Check" not in record.full_text
        assert "Degree in a relevant discipline" not in record.full_text
        excluded_reasons = {item["reason"] for item in record.excluded_items}
        assert "mandatory key accountability boilerplate" in excluded_reasons
        assert "mandatory essential requirement boilerplate" in excluded_reasons

        summary = generate_mapping_texts(connection)
        assert summary["generated"] == 1
        assert summary["excluded_items"] == len(record.excluded_items)
        stored = connection.execute(
            "SELECT full_text, excluded_items_json FROM pd_mapping_texts WHERE position_description_id = ?",
            (pd_id,),
        ).fetchone()
        assert "Primary purpose" in stored["full_text"]
        assert "mandatory key accountability boilerplate" in stored["excluded_items_json"]
    finally:
        connection.close()


def test_embedding_generation_stores_pd_and_framework_vectors(tmp_path):
    workbook_path = tmp_path / "mapping.xlsx"
    workbook = Workbook()
    mapping = workbook.active
    mapping.title = "job_mapping_input"
    mapping.append([
        "pd_id", "position_title", "position_grade",
        "primary_mapping_code", "primary_mapping_name", "primary_mapping_level",
        "secondary_mapping_code", "secondary_mapping_name", "secondary_mapping_level",
        "secondary_2_mapping_code", "secondary_2_mapping_name", "secondary_2_mapping_level",
        "mapping_validated", "validation_notes",
    ])
    mapping.append([10021, "Facilities Officer", "TW4", 2000000000, "FACILITIES", "Job Family", None, None, None, None, None, None, "Yes", None])
    framework = workbook.create_sheet("job_family_power_query")
    framework.append([
        "Code", "Name", "Level", "Parent Code", "main_definition",
        "supp_definition_1", "supp_definition_2", "exclusions",
    ])
    framework.append([
        2000000000,
        "FACILITIES",
        "Job Family",
        None,
        "Facilities roles manage physical places and services.",
        "Additional positive facilities context.",
        "More positive facilities context.",
        "Do not reward this exclusion text.",
    ])
    workbook.save(workbook_path)

    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        import_document(
            connection,
            extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx"),
        )
        import_job_family_workbook(connection, workbook_path)
        generate_mapping_texts(connection)

        summary = generate_embeddings(
            connection,
            FakeEmbedder(),
            model_name="fake-model",
            batch_size=2,
        )

        assert summary["pd_embeddings"] == 1
        assert summary["framework_embeddings"] == 1
        assert connection.execute("SELECT COUNT(*) FROM embeddings").fetchone()[0] == 2
        row = connection.execute(
            "SELECT embedding_dimension, source_text_hash, embedding_json FROM embeddings WHERE source_type = 'pd'"
        ).fetchone()
        assert row["embedding_dimension"] == 3
        assert len(row["source_text_hash"]) == 64
        assert row["embedding_json"].startswith("[")
        framework_row = connection.execute(
            "SELECT source_text FROM embeddings WHERE source_type = 'framework'"
        ).fetchone()
        assert "Additional positive facilities context" in framework_row["source_text"]
        assert "More positive facilities context" in framework_row["source_text"]
        assert "Do not reward this exclusion text" not in framework_row["source_text"]
    finally:
        connection.close()


def test_similar_role_search_ranks_pd_embeddings_and_includes_mapping(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        facilities = extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx")
        facilities_id = import_document(connection, facilities)
        cfo = extract_document(SAMPLES / "10773-02 Chief Financial Officer - PSSE Band 2.docx")
        cfo_id = import_document(connection, cfo)
        adr = extract_document(SAMPLES / "10108-01 Alternate Dispute Resolution Specialist - TW9.docx")
        adr_id = import_document(connection, adr)
        connection.executemany(
            """INSERT INTO embeddings(
                source_type, source_id, model_name, embedding_dimension,
                source_text_hash, source_text, embedding_json
            ) VALUES ('pd', ?, 'test-model', 3, 'hash', 'text', ?)""",
            [
                (str(facilities_id), "[1,0,0]"),
                (str(cfo_id), "[0.9,0.1,0]"),
                (str(adr_id), "[0,1,0]"),
            ],
        )
        connection.execute(
            """INSERT INTO job_family_import_batches(
                source_filename, framework_sheet, mapping_sheet, is_active
            ) VALUES ('test.xlsx', 'job_family_power_query', 'job_mapping_input', 1)"""
        )
        batch_id = connection.execute("SELECT id FROM job_family_import_batches").fetchone()["id"]
        connection.execute(
            """INSERT INTO pd_job_family_mappings(
                import_batch_id, position_description_id, pd_id, mapping_rank,
                mapping_code, mapping_name, mapping_level, mapping_validated, framework_code_valid
            ) VALUES (?, ?, '10773', 1, '2000', 'FINANCE', 'Job Family', 'Yes', 1)""",
            (batch_id, cfo_id),
        )
        connection.commit()

        results = find_similar_roles(
            connection,
            facilities_id,
            model_name="test-model",
            limit=2,
        )

        assert [result.position_description_id for result in results] == [cfo_id, adr_id]
        assert results[0].primary_mapping_name == "FINANCE"
        assert results[0].similarity > results[1].similarity
        assert cosine_similarity([1, 0, 0], [1, 0, 0]) == 1
    finally:
        connection.close()


def test_free_text_role_search_compares_query_to_pd_embeddings(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        facilities_id = import_document(
            connection,
            extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx"),
        )
        cfo_id = import_document(
            connection,
            extract_document(SAMPLES / "10773-02 Chief Financial Officer - PSSE Band 2.docx"),
        )
        connection.executemany(
            """INSERT INTO embeddings(
                source_type, source_id, model_name, embedding_dimension,
                source_text_hash, source_text, embedding_json
            ) VALUES ('pd', ?, 'test-model', 3, 'hash', 'text', ?)""",
            [
                (str(facilities_id), "[1,0,0]"),
                (str(cfo_id), "[0,1,0]"),
            ],
        )
        results = search_roles(
            connection,
            "find a facilities role",
            QueryFakeEmbedder(),
            model_name="test-model",
            limit=2,
        )
        assert [result.position_description_id for result in results] == [facilities_id, cfo_id]
        assert results[0].similarity == 1
    finally:
        connection.close()


def test_intelligence_app_pd_search_returns_mapping_context(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        pd_id = import_document(
            connection,
            extract_document(SAMPLES / "10021-01 Facilities Officer - TW4.docx"),
        )
        connection.execute(
            """INSERT INTO job_family_import_batches(
                source_filename, framework_sheet, mapping_sheet, is_active
            ) VALUES ('test.xlsx', 'job_family_power_query', 'job_mapping_input', 1)"""
        )
        batch_id = connection.execute("SELECT id FROM job_family_import_batches").fetchone()["id"]
        connection.execute(
            """INSERT INTO pd_job_family_mappings(
                import_batch_id, position_description_id, pd_id, mapping_rank,
                mapping_code, mapping_name, mapping_level, mapping_validated, framework_code_valid
            ) VALUES (?, ?, '10021', 1, '2000', 'FACILITIES', 'Job Family', 'Yes', 1)""",
            (batch_id, pd_id),
        )
        rows = search_pds(connection, "Facilities")
        assert rows[0]["role_title"] == "Facilities Officer"
        assert rows[0]["mapping_name"] == "FACILITIES"
    finally:
        connection.close()


def test_mapping_suggestions_group_similar_mapped_roles():
    roles = [
        SimilarRole(1, "10001-01", "Policy Officer", "a.docx", 0.90, "1021000000", "STRATEGIC POLICY", "Job Family", "Yes"),
        SimilarRole(2, "10002-01", "Senior Policy Officer", "b.docx", 0.85, "1021000000", "STRATEGIC POLICY", "Job Family", "Yes"),
        SimilarRole(3, "10003-01", "Policy Governance Lead", "c.docx", 0.80, "1010002000", "POLICY GOVERNANCE", "Sub-family", "Yes"),
        SimilarRole(4, "10004-01", "Unmapped Role", "d.docx", 0.95, "", "", "", ""),
        SimilarRole(5, "10005-01", "Director Policy", "e.docx", 0.78, "1021000000", "STRATEGIC POLICY", "Job Family", ""),
    ]

    suggestions = suggest_mappings_from_similar_roles(roles)

    assert suggestions[0].mapping_code == "1021000000"
    assert suggestions[0].evidence_count == 3
    assert suggestions[0].validated_evidence_count == 2
    assert suggestions[0].confidence == "High"
    assert suggestions[1].mapping_code == "1010002000"


def test_mapping_suggestion_hierarchy_includes_parent_levels(tmp_path):
    connection = connect_database(tmp_path / "test.sqlite3")
    initialise_database(connection)
    try:
        connection.execute(
            """INSERT INTO job_family_import_batches(
                source_filename, framework_sheet, mapping_sheet, is_active
            ) VALUES ('test.xlsx', 'job_family_power_query', 'job_mapping_input', 1)"""
        )
        batch_id = connection.execute("SELECT id FROM job_family_import_batches").fetchone()["id"]
        connection.executemany(
            """INSERT INTO job_family_entries(
                import_batch_id, code, name, level, parent_code, sequence
            ) VALUES (?, ?, ?, ?, ?, ?)""",
            [
                (batch_id, "1000000000", "CORPORATE", "Job Family", "", 1),
                (batch_id, "1000001000", "POLICY", "Sub-family", "1000000000", 2),
                (batch_id, "1000001001", "STRATEGIC POLICY", "Specialisation", "1000001000", 3),
            ],
        )
        suggestions = suggest_mappings_from_similar_roles([
            SimilarRole(1, "1", "Role", "a.docx", 0.9, "1000001001", "STRATEGIC POLICY", "Specialisation", "Yes")
        ])
        enrich_mapping_suggestion_hierarchies(connection, suggestions)
        assert [item["level"] for item in suggestions[0].hierarchy] == [
            "Job Family", "Sub-family", "Specialisation",
        ]
        assert [item["code"] for item in suggestions[0].hierarchy] == [
            "1000000000", "1000001000", "1000001001",
        ]
    finally:
        connection.close()
