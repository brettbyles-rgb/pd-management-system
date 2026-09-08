from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from pd_extractor import extract_document
from pd_extractor.database import connect_database, import_document, initialise_database
from pd_extractor.job_family_import import import_job_family_workbook
from pd_extractor.reference_workbook import (
    ADJACENCY_SHEET,
    FRAMEWORK_SHEET,
    MAPPING_SHEET,
    apply_reference_workbook,
    export_reference_workbook,
    validate_reference_workbook,
)


FIXTURES = Path(__file__).parent / "fixtures"


def _seed_database(tmp_path):
    connection = connect_database(tmp_path / "reference.sqlite3")
    initialise_database(connection)
    import_document(
        connection,
        extract_document(FIXTURES / "10021-01 Facilities Officer - TW4.docx"),
    )
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    mapping = workbook.active
    mapping.title = MAPPING_SHEET
    mapping.append(
        [
            "pd_id", "position_title", "position_grade",
            "primary_mapping_code", "primary_mapping_name", "primary_mapping_level",
            "secondary_mapping_code", "secondary_mapping_name", "secondary_mapping_level",
            "secondary_2_mapping_code", "secondary_2_mapping_name", "secondary_2_mapping_level",
            "mapping_validated", "validation_notes",
        ]
    )
    mapping.append([10021, "Facilities Officer", "TW4", 10, "Facilities", "Job Family", "", "", "", "", "", "", "Yes", "Reviewed"])
    framework = workbook.create_sheet(FRAMEWORK_SHEET)
    framework.append(["code", "name", "level", "parent_code", "main_definition"])
    framework.append([10, "Facilities", "Job Family", "", "Facilities work"])
    framework.append([20, "People", "Job Family", "", "People work"])
    workbook.save(source)
    workbook.close()
    summary = import_job_family_workbook(connection, source)
    connection.executemany(
        "INSERT INTO job_family_adjacency_entries VALUES (?, ?, ?, ?, ?)",
        [
            (summary.import_batch_id, "10", "10", 0, ""),
            (summary.import_batch_id, "10", "20", 1, ""),
            (summary.import_batch_id, "20", "10", 1, ""),
            (summary.import_batch_id, "20", "20", 0, ""),
        ],
    )
    connection.commit()
    return connection


def test_reference_workbook_round_trip_is_valid_and_versioned(tmp_path):
    connection = _seed_database(tmp_path)
    try:
        connection.execute(
            "UPDATE job_family_entries SET main_definition = '=literal text' WHERE name = 'People'"
        )
        connection.commit()
        exported = tmp_path / "governed-reference.xlsx"
        exported.write_bytes(export_reference_workbook(connection))
        workbook = load_workbook(exported, read_only=True, data_only=True)
        assert workbook.sheetnames == ["README", FRAMEWORK_SHEET, MAPPING_SHEET, ADJACENCY_SHEET]
        assert workbook[ADJACENCY_SHEET].max_row == 4
        people_definition = next(
            row[4].value
            for row in workbook[FRAMEWORK_SHEET].iter_rows(min_row=2)
            if row[1].value == "People"
        )
        assert people_definition == "=literal text"
        workbook.close()

        preview = validate_reference_workbook(connection, exported)
        assert preview["valid"] is True
        assert preview["counts"] == {
            "framework_rows": 2,
            "mapping_role_rows": 1,
            "mapping_assignments": 1,
            "job_families": 2,
            "adjacency_cells": 4,
        }
        assert all(value == 0 for value in preview["changes"].values())

        old_batch = connection.execute(
            "SELECT id FROM job_family_import_batches WHERE is_active = 1"
        ).fetchone()["id"]
        result = apply_reference_workbook(connection, exported)
        assert result["import_batch_id"] != old_batch
        assert connection.execute(
            "SELECT is_active FROM job_family_import_batches WHERE id = ?", (old_batch,)
        ).fetchone()["is_active"] == 0
        assert connection.execute(
            "SELECT COUNT(*) FROM active_job_family_entries"
        ).fetchone()[0] == 2
        assert connection.execute(
            "SELECT COUNT(*) FROM active_pd_job_family_mappings"
        ).fetchone()[0] == 1
        assert connection.execute(
            "SELECT COUNT(*) FROM active_job_family_adjacency"
        ).fetchone()[0] == 4
    finally:
        connection.close()


def test_reference_workbook_rejects_incomplete_adjacency_without_changes(tmp_path):
    connection = _seed_database(tmp_path)
    try:
        exported = tmp_path / "invalid-reference.xlsx"
        exported.write_bytes(export_reference_workbook(connection))
        workbook = load_workbook(exported)
        workbook[ADJACENCY_SHEET]["D3"] = None
        workbook.save(exported)
        workbook.close()

        active_before = connection.execute(
            "SELECT id FROM job_family_import_batches WHERE is_active = 1"
        ).fetchone()["id"]
        preview = validate_reference_workbook(connection, exported)
        assert preview["valid"] is False
        assert any("must be 0, 1 or 2" in error for error in preview["errors"])

        try:
            apply_reference_workbook(connection, exported)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid workbook was applied")
        active_after = connection.execute(
            "SELECT id FROM job_family_import_batches WHERE is_active = 1"
        ).fetchone()["id"]
        assert active_after == active_before
    finally:
        connection.close()
