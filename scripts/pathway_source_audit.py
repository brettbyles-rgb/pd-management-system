from __future__ import annotations

import csv
import json
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "output" / "pd-management-bulk-fixed.sqlite3"
WORKBOOK = (
    ROOT
    / "data"
    / "source"
    / "original"
    / "activity-assigned-clustered-role-links-revised-high-all.xlsx"
)


def query(connection: sqlite3.Connection, statement: str) -> list[dict[str, object]]:
    return [dict(row) for row in connection.execute(statement)]


def csv_profile(path: Path) -> dict[str, object]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
    id_field = "pd_id" if "pd_id" in (reader.fieldnames or []) else None
    return {
        "path": str(path.relative_to(ROOT)),
        "rows": len(rows),
        "fields": reader.fieldnames,
        "distinct_pd_ids": len({row[id_field] for row in rows}) if id_field else None,
        "sample": rows[:2],
    }


def main() -> int:
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    database = {
        "roles": query(
            connection,
            "SELECT COUNT(*) AS count, MIN(id) AS min_id, MAX(id) AS max_id "
            "FROM position_descriptions",
        )[0],
        "missing_pd_numbers": query(
            connection,
            "SELECT COUNT(*) AS count FROM position_descriptions "
            "WHERE TRIM(COALESCE(position_description_no, '')) = ''",
        )[0]["count"],
        "duplicate_pd_numbers": query(
            connection,
            "SELECT position_description_no, COUNT(*) AS records, "
            "COUNT(DISTINCT role_title) AS distinct_titles "
            "FROM position_descriptions "
            "WHERE TRIM(COALESCE(position_description_no, '')) <> '' "
            "GROUP BY position_description_no HAVING COUNT(*) > 1 "
            "ORDER BY records DESC, position_description_no",
        ),
        "duplicate_titles": query(
            connection,
            "SELECT role_title, COUNT(*) AS records FROM position_descriptions "
            "GROUP BY role_title HAVING COUNT(*) > 1 "
            "ORDER BY records DESC, role_title",
        ),
        "capability_frameworks": query(
            connection,
            "SELECT cf.framework_name, COUNT(DISTINCT cd.id) AS capabilities, "
            "COUNT(pc.id) AS assignments FROM capability_frameworks cf "
            "LEFT JOIN capability_definitions cd ON cd.framework_id = cf.id "
            "LEFT JOIN pd_capabilities pc ON pc.capability_definition_id = cd.id "
            "GROUP BY cf.id ORDER BY assignments DESC, cf.framework_name",
        ),
        "capability_levels": query(
            connection,
            "SELECT required_level, COUNT(*) AS assignments FROM pd_capabilities "
            "GROUP BY required_level ORDER BY assignments DESC, required_level",
        ),
        "roles_with_focus_capabilities": query(
            connection,
            "SELECT COUNT(DISTINCT position_description_id) AS count "
            "FROM pd_capabilities WHERE LOWER(capability_type) = 'focus'",
        )[0]["count"],
        "job_family_nodes": query(
            connection, "SELECT COUNT(*) AS count FROM active_job_family_entries"
        )[0]["count"],
        "role_job_family_mappings": query(
            connection,
            "SELECT COUNT(*) AS mappings, "
            "COUNT(DISTINCT position_description_id) AS roles "
            "FROM active_pd_job_family_mappings "
            "WHERE position_description_id IS NOT NULL",
        )[0],
        "capability_status_by_cohort": query(
            connection,
            "SELECT COALESCE(cr.cohort, '<unmapped>') AS cohort, COUNT(*) AS roles, "
            "SUM(CASE WHEN EXISTS (SELECT 1 FROM pd_capabilities pc "
            "  WHERE pc.position_description_id = pd.id) THEN 1 ELSE 0 END) "
            "  AS roles_with_capabilities, "
            "SUM(CASE WHEN EXISTS (SELECT 1 FROM pd_capabilities pc "
            "  WHERE pc.position_description_id = pd.id "
            "    AND LOWER(pc.capability_type) = 'focus') THEN 1 ELSE 0 END) "
            "  AS roles_with_focus "
            "FROM position_descriptions pd "
            "LEFT JOIN classification_references cr "
            "  ON cr.raw_label = pd.classification_grade_band "
            "GROUP BY COALESCE(cr.cohort, '<unmapped>') ORDER BY cohort",
        ),
    }
    connection.close()

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    workbook_profile: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        workbook_profile.append(
            {
                "sheet": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "fields": header,
            }
        )
    summary_sheet = workbook["Role Summary"]
    workbook_ids = [
        row[0]
        for row in summary_sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    workbook.close()

    csv_dir = ROOT / "data" / "source" / "onedrive_exports"
    csv_files = [csv_profile(path) for path in sorted(csv_dir.glob("*.csv"))]
    combined = json.loads(
        (
            ROOT
            / "data"
            / "source"
            / "original"
            / "pd-data-for-claude-career-pathways-with-classification-order.json"
        ).read_text(encoding="utf-8")
    )
    pd_numbers = [row.get("pd_id") for row in combined]
    titles = [str(row.get("title") or "") for row in combined]

    report = {
        "database": database,
        "workbook": {
            "sheets": workbook_profile,
            "role_ids": {
                "records": len(workbook_ids),
                "distinct": len(set(workbook_ids)),
                "minimum": min(workbook_ids),
                "maximum": max(workbook_ids),
            },
        },
        "csv_files": csv_files,
        "combined_json": {
            "records": len(combined),
            "fields": list(combined[0]) if combined else [],
            "missing_pd_numbers": sum(value in (None, "") for value in pd_numbers),
            "duplicate_pd_numbers": {
                str(key): count
                for key, count in Counter(pd_numbers).items()
                if key not in (None, "") and count > 1
            },
            "duplicate_titles": {
                key: count
                for key, count in Counter(titles).items()
                if key and count > 1
            },
        },
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
